"""Verbatim Vernehmlassungs corpus from Fedlex — Erläuternde Berichte,
Vernehmlassungsvorlagen, Ergebnisberichte, Stellungnahmen.

Fedlex publishes every federal consultation procedure (Vernehmlassung /
consultation / consultazione, 1993–today) as a ``jolux:Consultation``
(``eli/dl/proj/<y>/<n>/cons_<k>``). Its documents hang off the procedure's
sub-tasks and are typed by ``jolux:draftProcessDocumentType``::

    ?cons jolux:hasSubTask ?task .
    ?task jolux:opinionHasDraftRelatedDocument | jolux:opinionIsAboutDraftDocument ?doc .
    ?doc  jolux:draftProcessDocumentType <vocabulary/draft-document-type/N> ;
          jolux:isRealizedBy ?expr .                     # one per language
    ?expr jolux:language ?lang ; jolux:isEmbodiedBy ?m .  # pdf-a, docx, ...
    ?m    jolux:isExemplifiedBy ?file .                   # Fedlex file store

Census 2026-10-19: 2,581 procedures; 2,434 Erläuternde Berichte, 3,275
Vorlagen, 1,943 Ergebnisberichte, 1,214 Stellungnahmen collections, 243
synoptic tables; Erläuternde Berichte from 2005 on (earlier procedures are
mostly metadata without files on Fedlex).

The corpus lives in its own database, not in materialien.db: publish.py's
Step 2f copies every foreign table of materialien.db on each build and is a
fatal step with a time cap, and the Stellungnahmen alone are hundreds of
pages per procedure. Two files:

* the WORK database (``--work-db``, writable, never served) holds the
  discovery metadata, the per-file fetch state and the text;
* ``publish`` snapshots it with ``VACUUM INTO`` next to the real path of
  ``--out`` and ``os.replace``s it over the served file, so MCP workers
  (``?mode=ro&immutable=1``, one connection per call) never see a
  half-written database.

Documents Fedlex does not hold (the Chancellery's 1995-2004 index and the
BJ legislative projects, live and from the Wayback Machine) are discovered
into the same work DB by search_stack/consultation_archive.py and fetched
here like Fedlex's, each host at its own pace (HOST_MIN_INTERVAL); their
rows carry their provenance (source, original_url, source_page_url).

Text policy: the documents that carry legal reasoning or the draft itself are
ingested verbatim (Erläuternder Bericht, Ergebnisbericht, Vorlage, synoptic
table, Diverses, Stellungnahmen); cover letters, address lists, answer forms
and questionnaires are recorded with their Fedlex file link only.

Formats (2026-10-21): PDF, Word (.docx; legacy .doc through LibreOffice
headless), Excel (.xlsx row text per sheet; .xls through LibreOffice) and ZIP
archives (each PDF / Word / Excel member, its name kept on every paragraph as
``source_member``). The format is read from the bytes, not the label. Files
above 250 MB are streamed to disk. Scans without a text layer, and the
text-less pages of other PDFs, are read by the separate ``ocr`` pass
(Tesseract); OCR text is flagged ``text_source='ocr'`` per paragraph and is
never served as the document's own wording.

Usage::

    python3 -m search_stack.build_consultation_corpus discover --work-db W
    python3 -m search_stack.build_consultation_corpus fetch --work-db W \\
        --max-seconds 10800 --rate-limit 0.5
    python3 -m search_stack.build_consultation_corpus publish --work-db W \\
        --out output/consultations.db
    python3 -m search_stack.build_consultation_corpus status --work-db W
    python3 -m search_stack.build_consultation_corpus requeue --work-db W \\
        --what unsupported,too_large,parse_error,no_pages [--apply]
    python3 -m search_stack.build_consultation_corpus ocr --work-db W \\
        --scope no_text[,partial] --max-seconds 7200 --workers 2
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import logging
import os
import re
import sqlite3
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Iterable, Iterator

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from search_stack.build_botschaft_corpus import (  # noqa: E402
    ARTICLE_HEADER_RE,
    BACKOFF_BASE_SECONDS,
    BACKOFF_MAX_SECONDS,
    FATAL_HTTP_STATUSES,
    INGEST_STOPPED_EXIT_CODE,
    STOP_AFTER_CONSECUTIVE_FAILURES,
    SourceUnavailable,
    http_status_of,
)
from search_stack.botschaft_historical import _join_hyphen  # noqa: E402
from scrapers.fedlex_materialien import fix_cp1252_controls  # noqa: E402

log = logging.getLogger("build_consultation_corpus")

JOLUX = "http://data.legilux.public.lu/resource/ontology/jolux#"
SKOS = "http://www.w3.org/2004/02/skos/core#"
VOCAB = "https://fedlex.data.admin.ch/vocabulary/"
LANG_FROM_URI = {
    "http://publications.europa.eu/resource/authority/language/DEU": "de",
    "http://publications.europa.eu/resource/authority/language/FRA": "fr",
    "http://publications.europa.eu/resource/authority/language/ITA": "it",
    "http://publications.europa.eu/resource/authority/language/ROH": "rm",
    "http://publications.europa.eu/resource/authority/language/ENG": "en",
}

USER_AGENT = ("opencaselaw-materialien/0.4 (Vernehmlassungen; "
              "+https://opencaselaw.ch)")
FEDLEX_FILE_HOST = "fedlex.data.admin.ch"
# Fedlex's file store, plus the two sources of search_stack.consultation_archive:
# Wayback copies of the admin.ch and bj.admin.ch pages Fedlex never took over,
# and the live bj.admin.ch file store.
ALLOWED_FILE_HOSTS = frozenset({FEDLEX_FILE_HOST, "web.archive.org", "www.bj.admin.ch"})
# Minimum seconds after a request to a host before the next request of the
# run (one connection; the Internet Archive asks for >= 2 s between requests).
HOST_MIN_INTERVAL: dict[str, float] = {"web.archive.org": 2.0, "www.bj.admin.ch": 1.0}
# The Wayback Machine answers 403 for a single capture excluded from replay:
# that file is final (not_archived), only 429 is a refusal of the client.
WAYBACK_HOST = "web.archive.org"

# draft-document-type code → slug. Labels are never synthesised: they are
# read from Fedlex's vocabulary at discovery time (consultation_vocab).
DOC_TYPES: dict[int, str] = {
    11: "vernehmlassungsvorlage",
    12: "erlaeuternder_bericht",
    13: "begleitschreiben",
    14: "adressatenliste",
    15: "vernehmlassung_diverses",
    16: "synoptische_tabelle",
    21: "stellungnahmen",
    22: "antwortformular",
    23: "fragebogen",
    31: "ergebnisbericht",
}
# Ingest order: the documents most cited for legislative intent first, the
# bulky collections of submissions last.
TEXT_TYPE_PRIORITY: tuple[int, ...] = (12, 31, 11, 16, 15, 21)
TEXT_TYPES = frozenset(TEXT_TYPE_PRIORITY)

# Manifestation preference when one expression has several files.
_FORMAT_PRIORITY = ("pdf-a", "pdf", "docx", "doc", "xlsx", "zip")
# Legacy Word (.doc, converted by LibreOffice), Excel workbooks (row text)
# and ZIP archives (their PDF / Word / Excel members) since 2026-10-21; until
# then 567 Fedlex files were kept as unsupported_format.
TEXT_FORMATS = frozenset({"pdf-a", "pdf", "docx", "doc", "xlsx", "zip"})

MAX_FILE_BYTES = 250 * 1024 * 1024
# A file above MAX_FILE_BYTES is streamed to a temporary file next to the
# work DB instead of being held in memory (2026-10-21: eight Stellungnahmen
# collections of 270-512 MB, each filed under all three languages).
MAX_STREAM_BYTES = 1024 * 1024 * 1024
SAME_FILE_PREFIX_BYTES = 256 * 1024
# ZIP members: limits against a corrupt or hostile archive.
MAX_ZIP_MEMBERS = 500
MAX_ZIP_TOTAL_BYTES = 1024 * 1024 * 1024
MAX_ZIP_DEPTH = 2
# Common Crawl cuts every record at 5 MiB; the Wayback Machine replays such a
# capture as if it were the file (2026-10-21: 19 BJ Stellungnahmen PDFs).
WAYBACK_TRUNCATED_BYTES = 5 * 1024 * 1024
MAX_ATTEMPTS = 5
FIRST_YEAR = 1990

# Fetch-state values of consultation_files.status.
ST_PENDING = "pending"
ST_DONE = "done"                  # text stored in consultation_paragraphs
ST_SAME_AS = "same_as"            # byte-identical to another language's file
ST_NO_TEXT = "no_text"            # scan without a text layer (OCR backlog)
ST_LINK_ONLY = "link_only"        # type outside the text policy
ST_UNSUPPORTED = "unsupported_format"
ST_NOT_FOUND = "not_found"        # file store answered 404/410
ST_TOO_LARGE = "too_large"
ST_PARSE_ERROR = "parse_error"    # downloaded but unreadable (corrupt / encrypted)
ST_ERROR = "error"                # retried until MAX_ATTEMPTS
ST_NOT_ARCHIVED = "not_archived"  # linked by an archived page, never captured by the Wayback Machine
ST_EMPTY = "empty_file"           # the source serves a zero-byte file
ST_TRUNCATED = "truncated"        # the archived copy is cut short (Common Crawl 5 MiB record)
FINAL_STATUSES = frozenset({ST_DONE, ST_SAME_AS, ST_NO_TEXT, ST_LINK_ONLY,
                            ST_UNSUPPORTED, ST_NOT_FOUND, ST_TOO_LARGE,
                            ST_PARSE_ERROR, ST_NOT_ARCHIVED, ST_EMPTY, ST_TRUNCATED})

# consultation_files.text_source / consultation_paragraphs.text_source:
# NULL = the file's own text (PDF text layer, Word, Excel); 'ocr' = read by
# Tesseract from the page image of a scan; a file whose text is partly
# either is 'mixed'. OCR text is never presented as the document's own
# wording (R2): the server flags every OCR passage.
TEXT_SOURCE_OCR = "ocr"
TEXT_SOURCE_MIXED = "mixed"


# ── Schema ─────────────────────────────────────────────────────────────

SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS consultation_procedures (
    procedure_uri    TEXT PRIMARY KEY,       -- https://fedlex.data.admin.ch/eli/dl/proj/2023/45/cons_1
    event_id         TEXT NOT NULL,          -- proj/2023/45/cons_1
    year             INTEGER,
    status_uri       TEXT,
    title_de         TEXT, title_fr TEXT, title_it TEXT,
    description_de   TEXT, description_fr TEXT, description_it TEXT,
    start_date       TEXT,                   -- opening of the consultation phase
    end_date         TEXT,                   -- deadline for submissions
    institution_uri  TEXT,                   -- most specific of Level3/2/1
    discovered_at    TEXT NOT NULL,
    updated_at       TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS ix_cp_year ON consultation_procedures(year);

-- The acts a procedure foresees to change (jolux:foreseenImpactToLegalResource).
CREATE TABLE IF NOT EXISTS consultation_laws (
    procedure_uri    TEXT NOT NULL REFERENCES consultation_procedures(procedure_uri),
    work_uri         TEXT NOT NULL,
    sr_number        TEXT,
    PRIMARY KEY (procedure_uri, work_uri)
);
CREATE INDEX IF NOT EXISTS ix_cl_sr ON consultation_laws(sr_number);

CREATE TABLE IF NOT EXISTS consultation_documents (
    document_uri     TEXT PRIMARY KEY,       -- …/cons_1/doc_10
    procedure_uri    TEXT NOT NULL REFERENCES consultation_procedures(procedure_uri),
    doc_type_code    INTEGER NOT NULL,
    doc_type         TEXT NOT NULL,          -- DOC_TYPES slug
    title_de         TEXT, title_fr TEXT, title_it TEXT   -- Fedlex's own ("Erläuternder Bericht-2")
);
CREATE INDEX IF NOT EXISTS ix_cd_proc ON consultation_documents(procedure_uri);
CREATE INDEX IF NOT EXISTS ix_cd_type ON consultation_documents(doc_type);

CREATE TABLE IF NOT EXISTS consultation_files (
    file_id          INTEGER PRIMARY KEY AUTOINCREMENT,
    document_uri     TEXT NOT NULL REFERENCES consultation_documents(document_uri),
    language         TEXT NOT NULL,
    file_url         TEXT,
    format           TEXT,
    file_name        TEXT,                   -- the expression title, usually the upload's file name
    publication_date TEXT,
    alt_files_json   TEXT,                   -- other manifestations [{format, url}]
    text_policy      TEXT NOT NULL,          -- 'text' | 'link'
    status           TEXT NOT NULL,
    http_status      INTEGER,
    attempts         INTEGER NOT NULL DEFAULT 0,
    last_attempt_at  TEXT,
    error            TEXT,
    byte_size        INTEGER,
    sha256           TEXT,
    prefix_sha256    TEXT,                   -- sha-256 of the first SAME_FILE_PREFIX_BYTES
    same_as_file_id  INTEGER,
    page_count       INTEGER,
    char_count       INTEGER,
    empty_pages      INTEGER,
    ingested_at      TEXT,
    UNIQUE(document_uri, language)
);
CREATE INDEX IF NOT EXISTS ix_cf_status ON consultation_files(status);

CREATE TABLE IF NOT EXISTS consultation_paragraphs (
    paragraph_id     INTEGER PRIMARY KEY AUTOINCREMENT,
    file_id          INTEGER NOT NULL REFERENCES consultation_files(file_id) ON DELETE CASCADE,
    para_order       INTEGER NOT NULL,
    page_number      INTEGER,
    section_path     TEXT,
    article_anchor   TEXT,
    text             TEXT NOT NULL,
    text_length      INTEGER NOT NULL,
    UNIQUE(file_id, para_order)
);
CREATE INDEX IF NOT EXISTS ix_cpa_file ON consultation_paragraphs(file_id);

CREATE VIRTUAL TABLE IF NOT EXISTS consultation_paragraphs_fts USING fts5(
    text,
    section_path,
    article_anchor UNINDEXED,
    content='consultation_paragraphs',
    content_rowid='paragraph_id',
    tokenize='unicode61 remove_diacritics 1'
);
CREATE TRIGGER IF NOT EXISTS consultation_paragraphs_ai AFTER INSERT ON consultation_paragraphs BEGIN
    INSERT INTO consultation_paragraphs_fts(rowid, text, section_path, article_anchor)
    VALUES (new.paragraph_id, new.text, new.section_path, new.article_anchor);
END;
CREATE TRIGGER IF NOT EXISTS consultation_paragraphs_ad AFTER DELETE ON consultation_paragraphs BEGIN
    INSERT INTO consultation_paragraphs_fts(consultation_paragraphs_fts, rowid, text, section_path, article_anchor)
    VALUES ('delete', old.paragraph_id, old.text, old.section_path, old.article_anchor);
END;

-- Fedlex vocabulary labels (document types, status, institutions) in DE/FR/IT.
CREATE TABLE IF NOT EXISTS consultation_vocab (
    uri              TEXT PRIMARY KEY,
    label_de         TEXT, label_fr TEXT, label_it TEXT
);

CREATE TABLE IF NOT EXISTS meta (
    key              TEXT PRIMARY KEY,
    value            TEXT
);

-- One row per procedure entry read from a non-Fedlex source page
-- (search_stack.consultation_archive): what the page said, where it was read,
-- and how it was matched to a procedure. Kept for every entry, matched or
-- not, so a match can be audited and re-done.
CREATE TABLE IF NOT EXISTS consultation_archive_entries (
    entry_id         TEXT PRIMARY KEY,       -- <source>:<key>
    source           TEXT NOT NULL,
    year             INTEGER,
    title_de         TEXT, title_fr TEXT, title_it TEXT,
    opened           TEXT,                   -- as printed on the page (ISO date)
    deadline         TEXT,
    department       TEXT,
    page_urls_json   TEXT NOT NULL,          -- {lang: page URL (Wayback snapshot or live)}
    procedure_uri    TEXT,                   -- matched or created procedure; NULL when ambiguous
    match_method     TEXT NOT NULL,          -- fedlex_link | deadline_title | opened_title | opened_institution | title | document_date | bj_live_twin | created | ambiguous | no_documents | documents_dated
    match_score      REAL,
    documents        INTEGER NOT NULL DEFAULT 0,
    discovered_at    TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS ix_cae_proc ON consultation_archive_entries(procedure_uri);
"""

# Provenance of rows that do not come from Fedlex (2026-10-20). Added with
# ALTER TABLE so a work database from before keeps its rows; NULL = Fedlex.
#   source           admin_pc_wayback | bj | bj_wayback
#   original_url     the file's own URL on the source site (dead for Wayback rows)
#   source_page_url  the page that linked it (Wayback snapshot or live page)
# consultation_files.file_url is what is fetched and shown: the Wayback
# snapshot URL for archived files, the live URL otherwise.
ARCHIVE_COLUMNS: dict[str, tuple[str, ...]] = {
    "consultation_procedures": ("source", "source_url"),
    "consultation_documents": ("source",),
    "consultation_files": ("source", "original_url", "source_page_url"),
}


# How the text of a file was obtained (2026-10-21). Additive like
# ARCHIVE_COLUMNS; NULL everywhere on rows from before.
#   consultation_files.text_source      NULL | 'ocr' | 'mixed' (TEXT_SOURCE_*)
#   consultation_files.ocr_pages        pages whose text is OCR
#   consultation_files.ocr_engine       "tesseract 5.3.4 deu 300dpi"
#   consultation_files.ocr_at           when the OCR pass looked at the file (also when it found nothing)
#   consultation_files.extraction_json  {"format": sniffed format, "method": ..., "members": [...], "skipped": [...]}
#   consultation_paragraphs.text_source NULL | 'ocr'
#   consultation_paragraphs.source_member  the ZIP member or Excel sheet the paragraph is from
EXTRACTION_COLUMNS: dict[str, tuple[tuple[str, str], ...]] = {
    "consultation_files": (("text_source", "TEXT"), ("ocr_pages", "INTEGER"),
                           ("ocr_engine", "TEXT"), ("ocr_at", "TEXT"),
                           ("extraction_json", "TEXT")),
    "consultation_paragraphs": (("text_source", "TEXT"), ("source_member", "TEXT")),
}


def ensure_archive_columns(conn: sqlite3.Connection) -> None:
    for table, cols in ARCHIVE_COLUMNS.items():
        have = {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}
        for col in cols:
            if col not in have:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN {col} TEXT")
    for table, typed in EXTRACTION_COLUMNS.items():
        have = {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}
        for col, typ in typed:
            if col not in have:
                conn.execute(f"ALTER TABLE {table} ADD COLUMN {col} {typ}")


def connect_work_db(path: str | Path) -> sqlite3.Connection:
    conn = sqlite3.connect(str(path), timeout=60)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA synchronous=NORMAL")
    conn.execute("PRAGMA foreign_keys=ON")
    conn.executescript(SCHEMA_SQL)
    ensure_archive_columns(conn)
    conn.commit()
    return conn


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def set_meta(conn: sqlite3.Connection, key: str, value) -> None:
    if not isinstance(value, str):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    conn.execute("INSERT INTO meta(key, value) VALUES (?, ?) "
                 "ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, value))


# ── SPARQL ─────────────────────────────────────────────────────────────

def _default_sparql(query: str) -> list[dict]:
    from scrapers.fedlex_materialien import sparql_query
    return sparql_query(query, timeout=300)


def _v(row: dict, key: str) -> str | None:
    cell = row.get(key)
    return cell["value"] if cell else None


def _year_filter(year: int) -> str:
    return f'FILTER(STRSTARTS(STR(?id), "proj/{year}/"))'


def q_procedures(year: int) -> str:
    return f"""
PREFIX jolux: <{JOLUX}>
SELECT DISTINCT ?c ?id ?status ?title ?tl ?desc WHERE {{
  ?c a jolux:Consultation ; jolux:eventId ?id .
  {_year_filter(year)}
  OPTIONAL {{ ?c jolux:consultationStatus ?status }}
  OPTIONAL {{
    ?c jolux:eventTitle ?title BIND(LANG(?title) AS ?tl)
    OPTIONAL {{ ?c jolux:eventDescription ?desc FILTER(LANG(?desc) = LANG(?title)) }}
  }}
}}"""


def q_phases(year: int) -> str:
    return f"""
PREFIX jolux: <{JOLUX}>
SELECT DISTINCT ?c ?start ?end ?i1 ?i2 ?i3 WHERE {{
  ?c a jolux:Consultation ; jolux:eventId ?id ; jolux:hasSubTask ?p .
  {_year_filter(year)}
  ?p a jolux:ConsultationPhase .
  OPTIONAL {{ ?p jolux:eventStartDate ?start }}
  OPTIONAL {{ ?p jolux:eventEndDate ?end }}
  OPTIONAL {{ ?p jolux:institutionInChargeOfTheEvent ?i1 }}
  OPTIONAL {{ ?p jolux:institutionInChargeOfTheEventLevel2 ?i2 }}
  OPTIONAL {{ ?p jolux:institutionInChargeOfTheEventLevel3 ?i3 }}
}}"""


def q_laws(year: int) -> str:
    # SR numbers: taxonomy notation UNION historicalLegalId — totally revised
    # acts carry no historicalLegalId (memory: fedlex-work-identity-trap).
    return f"""
PREFIX jolux: <{JOLUX}>
PREFIX skos: <{SKOS}>
SELECT DISTINCT ?c ?w ?sr WHERE {{
  ?c a jolux:Consultation ; jolux:eventId ?id ; jolux:foreseenImpactToLegalResource ?w .
  {_year_filter(year)}
  OPTIONAL {{
    {{ ?w jolux:classifiedByTaxonomyEntry ?tax . ?tax skos:notation ?n BIND(STR(?n) AS ?sr) }}
    UNION
    {{ ?w jolux:historicalLegalId ?h BIND(STR(?h) AS ?sr) }}
  }}
}}"""


def q_documents(year: int) -> str:
    return f"""
PREFIX jolux: <{JOLUX}>
SELECT DISTINCT ?c ?doc ?dt ?dtitle ?dtl ?lang ?fname ?fmt ?file ?pd WHERE {{
  ?c a jolux:Consultation ; jolux:eventId ?id ; jolux:hasSubTask ?s .
  {_year_filter(year)}
  VALUES ?rel {{ jolux:opinionHasDraftRelatedDocument jolux:opinionIsAboutDraftDocument }}
  ?s ?rel ?doc .
  ?doc jolux:draftProcessDocumentType ?dt .
  OPTIONAL {{ ?doc jolux:title ?dtitle BIND(LANG(?dtitle) AS ?dtl) }}
  OPTIONAL {{
    ?doc jolux:isRealizedBy ?e . ?e jolux:language ?lang .
    OPTIONAL {{ ?e jolux:title ?fname }}
    OPTIONAL {{
      ?e jolux:isEmbodiedBy ?m . ?m jolux:isExemplifiedBy ?file .
      OPTIONAL {{ ?m jolux:userFormat ?fmt }}
      OPTIONAL {{ ?m jolux:publicationDate ?pd }}
    }}
  }}
}}"""


def q_vocab() -> str:
    return f"""
PREFIX skos: <{SKOS}>
SELECT DISTINCT ?s ?l ?lg WHERE {{
  ?s skos:prefLabel ?l BIND(LANG(?l) AS ?lg)
  FILTER(?lg IN ("de", "fr", "it"))
  FILTER(STRSTARTS(STR(?s), "{VOCAB}draft-document-type/")
      || STRSTARTS(STR(?s), "{VOCAB}consultation-status/")
      || STRSTARTS(STR(?s), "{VOCAB}legal-institution/"))
}}"""


def q_year_counts() -> str:
    return f"""
PREFIX jolux: <{JOLUX}>
SELECT ?y (COUNT(DISTINCT ?c) AS ?n) WHERE {{
  ?c a jolux:Consultation ; jolux:eventId ?id .
  BIND(STR(SUBSTR(STR(?id), 6, 4)) AS ?y)
}} GROUP BY ?y ORDER BY ?y"""


def _format_of(fmt_uri: str | None, file_url: str | None) -> str | None:
    if fmt_uri:
        return fmt_uri.rstrip("/").rsplit("/", 1)[-1]
    if file_url and "." in file_url.rsplit("/", 1)[-1]:
        return file_url.rsplit(".", 1)[-1].lower()
    return None


def _pick_manifestation(files: list[dict]) -> tuple[dict | None, list[dict]]:
    """Choose one file per expression by _FORMAT_PRIORITY; the rest are kept
    as alternatives."""
    if not files:
        return None, []

    def rank(f):
        fmt = f.get("format") or ""
        return (_FORMAT_PRIORITY.index(fmt) if fmt in _FORMAT_PRIORITY
                else len(_FORMAT_PRIORITY), f.get("url") or "")

    ordered = sorted(files, key=rank)
    return ordered[0], ordered[1:]


def discover_years(sparql: Callable[[str], list[dict]]) -> list[int]:
    years = []
    for r in sparql(q_year_counts()):
        y = _v(r, "y")
        if y and y.isdigit() and int(y) >= FIRST_YEAR:
            years.append(int(y))
    return sorted(set(years))


def discover(conn: sqlite3.Connection, *,
             sparql: Callable[[str], list[dict]] = _default_sparql,
             years: Iterable[int] | None = None,
             pause: float = 0.5,
             sleep: Callable[[float], None] = time.sleep) -> dict:
    """Upsert procedures, laws, documents and file rows from Fedlex.

    Idempotent. Never deletes: a document that vanished from Fedlex keeps its
    row (and text) so a takedown is an explicit act, not a side effect of a
    flaky SPARQL answer. Existing file rows keep their fetch state unless the
    file URL changed, in which case they are re-queued."""
    now = _now()
    stats = {"years": 0, "procedures": 0, "documents": 0, "files": 0,
             "files_new": 0, "files_requeued": 0, "vocab": 0}

    for r in sparql(q_vocab()):
        uri, label, lg = _v(r, "s"), _v(r, "l"), _v(r, "lg")
        if not (uri and label and lg in ("de", "fr", "it")):
            continue
        conn.execute(
            f"INSERT INTO consultation_vocab(uri, label_{lg}) VALUES (?, ?) "
            f"ON CONFLICT(uri) DO UPDATE SET label_{lg}=excluded.label_{lg}",
            (uri, label))
        stats["vocab"] += 1
    conn.commit()

    for year in (sorted(years) if years else discover_years(sparql)):
        stats["years"] += 1
        procs: dict[str, dict] = {}
        for r in sparql(q_procedures(year)):
            c = _v(r, "c")
            if not c:
                continue
            p = procs.setdefault(c, {"event_id": _v(r, "id"), "status": None,
                                     "title": {}, "desc": {}})
            p["status"] = p["status"] or _v(r, "status")
            tl = _v(r, "tl")
            if tl in ("de", "fr", "it"):
                if _v(r, "title"):
                    p["title"].setdefault(tl, fix_cp1252_controls(_v(r, "title")))
                if _v(r, "desc"):
                    p["desc"].setdefault(tl, fix_cp1252_controls(_v(r, "desc")))
        sleep(pause)
        for r in sparql(q_phases(year)):
            p = procs.get(_v(r, "c") or "")
            if p is None:
                continue
            # Several phases (re-opened consultations): the earliest start
            # and the latest deadline describe the procedure.
            st, en = _v(r, "start"), _v(r, "end")
            if st and (not p.get("start") or st < p["start"]):
                p["start"] = st
            if en and (not p.get("end") or en > p["end"]):
                p["end"] = en
            inst = _v(r, "i3") or _v(r, "i2") or _v(r, "i1")
            p.setdefault("inst", inst)
        sleep(pause)
        for uri, p in procs.items():
            conn.execute(
                """INSERT INTO consultation_procedures
                   (procedure_uri, event_id, year, status_uri,
                    title_de, title_fr, title_it,
                    description_de, description_fr, description_it,
                    start_date, end_date, institution_uri, discovered_at, updated_at)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(procedure_uri) DO UPDATE SET
                     event_id=excluded.event_id, year=excluded.year,
                     status_uri=excluded.status_uri,
                     title_de=excluded.title_de, title_fr=excluded.title_fr,
                     title_it=excluded.title_it,
                     description_de=excluded.description_de,
                     description_fr=excluded.description_fr,
                     description_it=excluded.description_it,
                     start_date=excluded.start_date, end_date=excluded.end_date,
                     institution_uri=excluded.institution_uri,
                     updated_at=excluded.updated_at""",
                (uri, p["event_id"] or "", year, p["status"],
                 p["title"].get("de"), p["title"].get("fr"), p["title"].get("it"),
                 p["desc"].get("de"), p["desc"].get("fr"), p["desc"].get("it"),
                 p.get("start"), p.get("end"), p.get("inst"), now, now))
        stats["procedures"] += len(procs)

        for r in sparql(q_laws(year)):
            c, w = _v(r, "c"), _v(r, "w")
            if c not in procs or not w:
                continue
            sr = _v(r, "sr")
            conn.execute(
                "INSERT INTO consultation_laws(procedure_uri, work_uri, sr_number) "
                "VALUES (?,?,?) ON CONFLICT(procedure_uri, work_uri) DO UPDATE SET "
                "sr_number=COALESCE(consultation_laws.sr_number, excluded.sr_number)",
                (c, w, sr))
        sleep(pause)

        docs: dict[str, dict] = {}
        for r in sparql(q_documents(year)):
            c, doc, dt = _v(r, "c"), _v(r, "doc"), _v(r, "dt")
            if c not in procs or not doc or not dt:
                continue
            code_s = dt.rstrip("/").rsplit("/", 1)[-1]
            if not code_s.isdigit():
                continue
            d = docs.setdefault(doc, {"proc": c, "code": int(code_s),
                                      "title": {}, "files": {}})
            dtl = _v(r, "dtl")
            if dtl in ("de", "fr", "it") and _v(r, "dtitle"):
                d["title"].setdefault(dtl, _v(r, "dtitle"))
            lang = LANG_FROM_URI.get(_v(r, "lang") or "")
            if not lang:
                continue
            f = d["files"].setdefault(lang, {"name": None, "pd": None, "files": {}})
            f["name"] = f["name"] or _v(r, "fname")
            f["pd"] = f["pd"] or _v(r, "pd")
            url = _v(r, "file")
            if url:
                f["files"][url] = {"url": url, "format": _format_of(_v(r, "fmt"), url)}
        sleep(pause)

        for doc_uri, d in docs.items():
            code = d["code"]
            slug = DOC_TYPES.get(code, f"type_{code}")
            conn.execute(
                """INSERT INTO consultation_documents
                   (document_uri, procedure_uri, doc_type_code, doc_type,
                    title_de, title_fr, title_it)
                   VALUES (?,?,?,?,?,?,?)
                   ON CONFLICT(document_uri) DO UPDATE SET
                     procedure_uri=excluded.procedure_uri,
                     doc_type_code=excluded.doc_type_code, doc_type=excluded.doc_type,
                     title_de=excluded.title_de, title_fr=excluded.title_fr,
                     title_it=excluded.title_it""",
                (doc_uri, d["proc"], code, slug, d["title"].get("de"),
                 d["title"].get("fr"), d["title"].get("it")))
            stats["documents"] += 1
            for lang, f in d["files"].items():
                chosen, alts = _pick_manifestation(list(f["files"].values()))
                url = chosen["url"] if chosen else None
                fmt = chosen["format"] if chosen else None
                if code not in TEXT_TYPES:
                    policy, status = "link", ST_LINK_ONLY
                elif not url:
                    policy, status = "text", ST_NOT_FOUND
                elif fmt not in TEXT_FORMATS:
                    policy, status = "text", ST_UNSUPPORTED
                else:
                    policy, status = "text", ST_PENDING
                alt_json = json.dumps(alts, ensure_ascii=False) if alts else None
                prev = conn.execute(
                    "SELECT file_id, file_url, status FROM consultation_files "
                    "WHERE document_uri=? AND language=?", (doc_uri, lang)).fetchone()
                stats["files"] += 1
                if prev is None:
                    conn.execute(
                        """INSERT INTO consultation_files
                           (document_uri, language, file_url, format, file_name,
                            publication_date, alt_files_json, text_policy, status)
                           VALUES (?,?,?,?,?,?,?,?,?)""",
                        (doc_uri, lang, url, fmt, f["name"], f["pd"], alt_json,
                         policy, status))
                    stats["files_new"] += 1
                elif prev["file_url"] != url:
                    # Replaced upload: drop the old text, fetch again.
                    conn.execute("DELETE FROM consultation_paragraphs WHERE file_id=?",
                                 (prev["file_id"],))
                    conn.execute(
                        """UPDATE consultation_files SET file_url=?, format=?,
                           file_name=?, publication_date=?, alt_files_json=?,
                           text_policy=?, status=?, attempts=0, error=NULL,
                           http_status=NULL, sha256=NULL, prefix_sha256=NULL,
                           byte_size=NULL, same_as_file_id=NULL, page_count=NULL,
                           char_count=NULL, empty_pages=NULL, ingested_at=NULL
                           WHERE file_id=?""",
                        (url, fmt, f["name"], f["pd"], alt_json, policy, status,
                         prev["file_id"]))
                    stats["files_requeued"] += 1
                else:
                    conn.execute(
                        "UPDATE consultation_files SET file_name=?, publication_date=?, "
                        "alt_files_json=? WHERE file_id=?",
                        (f["name"], f["pd"], alt_json, prev["file_id"]))
        conn.commit()
        log.info("discover %d: %d procedures, %d documents", year, len(procs), len(docs))
    set_meta(conn, "last_discovery", {"at": now, **stats})
    conn.commit()
    return stats


# ── Text extraction ────────────────────────────────────────────────────

_PAGE_NUMBER_RE = re.compile(r"^\s*(?:Seite\s+|Page\s+|Pagina\s+)?\d{1,4}\s*(?:/\s*\d{1,4}|von\s+\d{1,4}|sur\s+\d{1,4}|di\s+\d{1,4})?\s*$",
                             re.IGNORECASE)
_SOFT_HYPHEN = "­"


def _norm_block(text: str) -> str:
    return re.sub(r"\d+", "#", re.sub(r"\s+", " ", text)).strip().lower()


def _join_lines(raw: str, language: str) -> str:
    lines = [ln.strip() for ln in raw.replace(_SOFT_HYPHEN, "").replace(" ", " ").split("\n")]
    lines = [ln for ln in lines if ln]
    if not lines:
        return ""
    out = lines[0]
    for nxt in lines[1:]:
        out = _join_hyphen(out, nxt, language)
    return re.sub(r"[ \t]{2,}", " ", out).strip()


def _open_pdf(src: bytes | str | Path):
    import fitz  # PyMuPDF
    if isinstance(src, (bytes, bytearray)):
        return fitz.open(stream=src, filetype="pdf")
    return fitz.open(str(src), filetype="pdf")


def pdf_page_blocks(src: bytes | str | Path) -> list[list[str]]:
    """One list of raw text blocks per page (PyMuPDF block order). ``src``
    is the PDF's bytes or the path of a streamed download."""
    pages: list[list[str]] = []
    with _open_pdf(src) as doc:
        for page in doc:
            try:
                blocks = page.get_text("blocks", sort=False)
            except Exception as e:  # noqa: BLE001 - a broken page must not lose the document
                log.warning("    page %d extract failed: %s", page.number + 1, e)
                blocks = []
            pages.append([b[4] for b in blocks if len(b) > 6 and b[6] == 0 and b[4].strip()])
    return pages


def docx_page_blocks(docx_bytes: bytes) -> list[list[str]]:
    """A .docx has no pages: one pseudo-page of its paragraphs (and table
    cells, row by row). python-docx where installed; otherwise the same
    reading of word/document.xml with the standard library (a host without
    python-docx once recorded three Ergebnisberichte as parse_error)."""
    try:
        import docx  # python-docx
    except ImportError:
        return _docx_page_blocks_xml(docx_bytes)
    d = docx.Document(io.BytesIO(docx_bytes))
    blocks = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(dict.fromkeys(cells)))
    return [blocks]


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _w_paragraph_text(p) -> str:
    """Text of a w:p as python-docx reads it: runs' w:t, w:tab as a tab,
    w:br / w:cr as a line break."""
    out = []
    for el in p.iter():
        if el.tag == _W + "t" and el.text:
            out.append(el.text)
        elif el.tag == _W + "tab":
            out.append("\t")
        elif el.tag in (_W + "br", _W + "cr"):
            out.append("\n")
    return "".join(out)


def _docx_page_blocks_xml(docx_bytes: bytes) -> list[list[str]]:
    """python-docx's order: the body's paragraphs, then each table row by
    row (cell text = its paragraphs joined by line breaks)."""
    import xml.etree.ElementTree as ET
    import zipfile
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    body = root.find(_W + "body")
    if body is None:
        return [[]]
    blocks = [t for t in (_w_paragraph_text(p) for p in body.findall(_W + "p")) if t.strip()]
    for tbl in body.iter(_W + "tbl"):
        for tr in tbl.findall(_W + "tr"):
            cells = []
            for tc in tr.findall(_W + "tc"):
                txt = "\n".join(_w_paragraph_text(p) for p in tc.findall(_W + "p")).strip()
                if txt:
                    cells.append(txt)
            if cells:
                blocks.append(" | ".join(dict.fromkeys(cells)))
    return [blocks]


_X = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def xlsx_sheet_blocks(xlsx_bytes: bytes) -> list[tuple[str, list[str]]]:
    """(sheet name, one block per non-empty row) for each worksheet, in
    workbook order. A row's cells are joined with " | " as table rows of a
    .docx are; values are the stored text (shared strings, inline strings,
    the cached value of a formula, numbers as stored). Standard library
    only: no openpyxl on the production host."""
    import posixpath
    import xml.etree.ElementTree as ET
    import zipfile
    with zipfile.ZipFile(io.BytesIO(xlsx_bytes)) as z:
        names = set(z.namelist())
        shared: list[str] = []
        if "xl/sharedStrings.xml" in names:
            for si in ET.fromstring(z.read("xl/sharedStrings.xml")).findall(_X + "si"):
                # rich text: the runs' w:t in order; phonetic runs (rPh) are not text
                shared.append("".join(t.text or "" for t in si.iter(_X + "t")
                                      if not _in_phonetic(si, t)))
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        targets = {}
        if "xl/_rels/workbook.xml.rels" in names:
            for rel in ET.fromstring(z.read("xl/_rels/workbook.xml.rels")).findall(_PKG_REL + "Relationship"):
                t = rel.get("Target") or ""
                targets[rel.get("Id")] = (t.lstrip("/") if t.startswith("/")
                                          else posixpath.normpath(posixpath.join("xl", t)))
        out: list[tuple[str, list[str]]] = []
        sheets = wb.find(_X + "sheets")
        for i, sh in enumerate(sheets.findall(_X + "sheet") if sheets is not None else [], 1):
            part = targets.get(sh.get(_R_ID)) or f"xl/worksheets/sheet{i}.xml"
            if part not in names:
                continue
            rows: list[str] = []
            for row in ET.fromstring(z.read(part)).iter(_X + "row"):
                cells = []
                for c in row.findall(_X + "c"):
                    v = _xlsx_cell_text(c, shared)
                    if v and v.strip():
                        cells.append(v.strip())
                if cells:
                    rows.append(" | ".join(cells))
            out.append((sh.get("name") or f"Sheet{i}", rows))
    return out


def _in_phonetic(si, t) -> bool:
    for rph in si.findall(_X + "rPh"):
        if any(x is t for x in rph.iter(_X + "t")):
            return True
    return False


def _xlsx_cell_text(c, shared: list[str]) -> str | None:
    typ = c.get("t")
    if typ == "inlineStr":
        is_ = c.find(_X + "is")
        return "".join(t.text or "" for t in is_.iter(_X + "t")) if is_ is not None else None
    v = c.find(_X + "v")
    if v is None or v.text is None:
        return None
    if typ == "s":
        try:
            return shared[int(v.text)]
        except (ValueError, IndexError):
            return None
    if typ == "b":
        return "TRUE" if v.text.strip() == "1" else "FALSE"
    return v.text


# ── legacy Office formats (LibreOffice) ───────────────────────────────

SOFFICE_TIMEOUT_SECONDS = 180


def soffice_convert(data: bytes, src_ext: str, target: str) -> bytes:
    """Convert legacy Office bytes (.doc, .xls) with LibreOffice headless,
    already installed on the production host; a private profile directory
    per call, so a stuck or concurrent soffice never blocks this one.
    Raises RuntimeError when soffice is missing or the conversion fails."""
    import shutil
    import subprocess
    import tempfile
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("LibreOffice (soffice) not installed")
    with tempfile.TemporaryDirectory(prefix="ocl-soffice-") as tmp:
        src = Path(tmp) / f"in.{src_ext}"
        src.write_bytes(data)
        profile = Path(tmp) / "profile"
        out_dir = Path(tmp) / "out"
        try:
            res = subprocess.run(
                [exe, f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
                 "--convert-to", target, "--outdir", str(out_dir), str(src)],
                capture_output=True, timeout=SOFFICE_TIMEOUT_SECONDS)
        except subprocess.TimeoutExpired as e:
            raise RuntimeError(f"soffice timed out after {SOFFICE_TIMEOUT_SECONDS} s") from e
        out = out_dir / f"in.{target}"
        if res.returncode != 0 or not out.exists():
            raise RuntimeError(f"soffice {src_ext}->{target} failed (rc={res.returncode}): "
                               f"{(res.stderr or b'').decode('utf-8', 'replace')[:200]}")
        return out.read_bytes()


# Replaced in tests (no LibreOffice on CI or the dev machines).
OFFICE_CONVERTER: Callable[[bytes, str, str], bytes] = soffice_convert


# ── format sniffing and extraction ─────────────────────────────────────

_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def sniff_format(head: bytes, declared: str | None = None, *,
                 whole: bytes | None = None) -> str | None:
    """The file's real format from its bytes: 'empty', 'pdf', 'docx',
    'xlsx', 'pptx', 'zip', 'doc', 'xls', 'ppt', else the declared one
    ('pdf-a' read as 'pdf'). The BJ file store serves .doc files under a
    PDF link, Fedlex labels a manifestation by its upload. ``whole`` is the
    complete file when it is in memory (an OLE file names its streams in a
    directory that can sit anywhere)."""
    if not head:
        return "empty"
    if head[:5] == b"%PDF-":
        return "pdf"
    if head[:8] == _OLE_MAGIC:
        data = whole if whole is not None else head
        if "WordDocument".encode("utf-16-le") in data:
            return "doc"
        if "Workbook".encode("utf-16-le") in data or "Book".encode("utf-16-le") in data:
            return "xls"
        if "PowerPoint Document".encode("utf-16-le") in data:
            return "ppt"
        return "ole"
    if head[:4] == b"PK\x03\x04" and whole is not None:
        import zipfile
        try:
            with zipfile.ZipFile(io.BytesIO(whole)) as z:
                names = set(z.namelist())
        except zipfile.BadZipFile:
            return "zip"
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "ppt/presentation.xml" in names:
            return "pptx"
        return "zip"
    if head[:4] == b"PK\x03\x04":
        return declared if declared in ("docx", "xlsx", "zip") else "zip"
    # a PDF after a few bytes of junk (a BOM, a mail header) is still a PDF
    if b"%PDF-" in head[:1024]:
        return "pdf"
    if declared == "pdf-a":
        return "pdf"
    return declared


class _Unsupported(Exception):
    """A format with no text extractor (PowerPoint, an unknown OLE file)."""


class TextPart:
    """One unit of text of a file: the file itself, or one ZIP member / PDF
    attachment / Excel sheet (``member``). ``src`` is the PDF of a paginated
    part (bytes or path), kept for the OCR pass."""
    __slots__ = ("member", "pages", "paginated", "ocr_pages", "src")

    def __init__(self, member: str | None, pages: list[list[str]], paginated: bool,
                 src: bytes | str | Path | None = None):
        self.member = member
        self.pages = pages
        self.paginated = paginated
        self.src = src
        self.ocr_pages: set[int] = set()   # 1-based page numbers whose blocks are OCR


def _zip_member_name(info) -> str:
    """Member name as the creator wrote it: UTF-8 when the archive says so
    (flag bit 11), else bytes the zipfile module decoded as cp437, which
    Windows tools write in the OEM code page (cp850: ä, ö, ü, é, è)."""
    if info.flag_bits & 0x800:
        return info.filename
    raw = info.filename.encode("cp437", errors="replace")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp850", errors="replace")


_ZIP_SKIP_RE = re.compile(r"(^|/)(__MACOSX/|\._|Thumbs\.db$|\.DS_Store$|desktop\.ini$)", re.IGNORECASE)


def extract_parts(src: bytes | str | Path, fmt: str | None, *, depth: int = 0,
                  prefix: str = "") -> tuple[list[TextPart], dict]:
    """Text parts of a file and how they were read (``extraction_json``).

    ``src`` is the file's bytes, or the path of a streamed PDF. Raises
    _Unsupported for a format without an extractor, any other exception for
    a file that cannot be read (the caller records parse_error)."""
    if not isinstance(src, (bytes, bytearray)):
        with open(src, "rb") as fh:
            head = fh.read(1024)
        real = sniff_format(head, fmt)
        if real == "pdf":
            return _pdf_parts(src, {"format": "pdf"}, depth, prefix)
        if head[:4] == b"PK\x03\x04":
            # a streamed archive is read member by member from disk
            return _zip_parts(src, {"format": "zip"}, depth, prefix)
        raise _Unsupported(f"streamed file is {real}, not a PDF or ZIP")
    data = bytes(src)
    real = sniff_format(data[:1024], fmt, whole=data)
    info: dict = {"format": real}
    if real == "pdf":
        return _pdf_parts(data, info, depth, prefix)
    if real == "docx":
        return [TextPart(prefix or None, docx_page_blocks(data), False)], info
    if real == "doc":
        info["method"] = "libreoffice doc->docx"
        return [TextPart(prefix or None, docx_page_blocks(OFFICE_CONVERTER(data, "doc", "docx")),
                         False)], info
    if real in ("xlsx", "xls"):
        if real == "xls":
            info["method"] = "libreoffice xls->xlsx"
            data = OFFICE_CONVERTER(data, "xls", "xlsx")
        parts = [TextPart(f"{prefix}{name}" if prefix else name, [rows], False)
                 for name, rows in xlsx_sheet_blocks(data)]
        info["sheets"] = [p.member for p in parts]
        return parts, info
    if real == "zip":
        return _zip_parts(io.BytesIO(data), info, depth, prefix)
    raise _Unsupported(f"no text extractor for {real or 'unknown'} files")


def _zip_parts(zsrc, info: dict, depth: int, prefix: str) -> tuple[list[TextPart], dict]:
    """The text parts of a ZIP archive's PDF / Word / Excel members (``zsrc``:
    a path or a file object), each named by its member path."""
    if depth >= MAX_ZIP_DEPTH:
        raise _Unsupported("ZIP nested too deep")
    import zipfile
    with zipfile.ZipFile(zsrc) as z:
        infos = [i for i in z.infolist() if not i.is_dir()]
        if len(infos) > MAX_ZIP_MEMBERS:
            raise ValueError(f"ZIP has {len(infos)} members > {MAX_ZIP_MEMBERS}")
        if sum(i.file_size for i in infos) > MAX_ZIP_TOTAL_BYTES:
            raise ValueError("ZIP expands beyond MAX_ZIP_TOTAL_BYTES")
        items = [(_zip_member_name(i), i.file_size, (lambda i=i: z.read(i))) for i in infos]
        parts = _member_parts(items, info, depth, prefix)
    return parts, info


def _member_parts(items: list, info: dict, depth: int, prefix: str) -> list[TextPart]:
    """Parts of the members of a container (ZIP archive, PDF attachments):
    ``items`` = (name, size, read()) in any order; read in name order, each
    part named by the member path; unreadable or unsupported members are
    listed in info['skipped'], never fatal."""
    parts: list[TextPart] = []
    members, skipped = [], []
    for name, size, read in sorted(items, key=lambda t: t[0]):
        if _ZIP_SKIP_RE.search(name):
            continue
        if size > MAX_FILE_BYTES:
            skipped.append({"member": name, "reason": f"{size} bytes"})
            continue
        ext = name.rsplit(".", 1)[-1].lower() if "." in name else None
        try:
            sub, sub_info = extract_parts(read(), ext, depth=depth + 1, prefix=f"{prefix}{name}/")
        except _Unsupported as e:
            skipped.append({"member": name, "reason": str(e)})
            continue
        except Exception as e:  # noqa: BLE001 - one broken member, not the container
            skipped.append({"member": name, "reason": f"unreadable: {e}"[:200]})
            continue
        for p in sub:
            # a member that is itself one part is named by the member
            if p.member == f"{prefix}{name}/":
                p.member = f"{prefix}{name}"
        parts.extend(sub)
        members.append({"member": name, "format": sub_info.get("format")})
    info["members"] = info.get("members", []) + members
    if skipped:
        info["skipped"] = info.get("skipped", []) + skipped
    return parts


def _fix_mojibake(name: str) -> str:
    """UTF-8 file names read as Latin-1 ("SantÃ© et addiction")."""
    if any(ch in name for ch in "ÃÂâ"):
        try:
            return name.encode("latin-1").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return name


def pdf_attachments(src: bytes | str | Path) -> tuple[bool, list[tuple[str, bytes]]]:
    """(is a portfolio, [(file name, bytes)]) of the files embedded in a PDF.
    Every file specification with an embedded stream is read: the Fedlex
    Stellungnahmen portfolios (2026-10-21) carry 9-101 submissions whose
    name tree MuPDF's embfile API does not always list."""
    out: list[tuple[str, bytes]] = []
    seen_streams: set[int] = set()
    seen_data: set[str] = set()

    def add(name: str | None, data: bytes | None, fallback: str) -> None:
        if not data:
            return
        h = hashlib.sha256(data).hexdigest()
        if h in seen_data:
            return
        seen_data.add(h)
        out.append((_fix_mojibake(name or fallback), data))

    with _open_pdf(src) as doc:
        cat = doc.pdf_catalog()
        portfolio = doc.xref_get_key(cat, "Collection")[0] not in ("null", "")
        # 1. the document's EmbeddedFiles name tree (inline file specifications)
        for i in range(doc.embfile_count()):
            try:
                inf = doc.embfile_info(i)
                add(inf.get("ufilename") or inf.get("filename") or inf.get("name"),
                    doc.embfile_get(i), f"attachment-{i + 1}")
            except Exception as e:  # noqa: BLE001 - one broken attachment, not the file
                log.warning("    attachment %d unreadable: %s", i, e)
        # 2. every file specification object with an embedded stream (its
        #    /Type /Filespec is optional; the Fedlex portfolios' name tree is
        #    not always listed by the API above)
        for x in range(1, doc.xref_length()):
            try:
                ef = doc.xref_get_key(x, "EF/F")
                if ef[0] != "xref":
                    continue
                sx = int(ef[1].split()[0])
                if sx in seen_streams:
                    continue
                seen_streams.add(sx)
                name = None
                for key in ("UF", "F"):
                    t, v = doc.xref_get_key(x, key)
                    if t == "string" and v:
                        name = v
                        break
                add(name, doc.xref_stream(sx), f"attachment-{sx}")
            except Exception as e:  # noqa: BLE001
                log.warning("    attachment %d unreadable: %s", x, e)
    return portfolio, out


def _pdf_parts(src: bytes | str | Path, info: dict, depth: int,
               prefix: str) -> tuple[list[TextPart], dict]:
    """A PDF's own pages, plus the files embedded in it as members. A PDF
    Portfolio's own page is the viewer's placeholder ("Zur optimalen Anzeige
    dieses PDF-Portfolios ..."), not the document: only its members count."""
    main = TextPart(prefix or None, pdf_page_blocks(src), True, src)
    portfolio, attached = (False, []) if depth >= MAX_ZIP_DEPTH else pdf_attachments(src)
    if not attached:
        return [main], info
    if len(attached) > MAX_ZIP_MEMBERS:
        raise ValueError(f"PDF has {len(attached)} attachments > {MAX_ZIP_MEMBERS}")
    info["attachments"] = len(attached)
    if portfolio:
        info["portfolio"] = True
    items = [(n, len(d), (lambda d=d: d)) for n, d in attached]
    members = _member_parts(items, info, depth, prefix)
    return ([] if portfolio else [main]) + members, info


_TOC_ENTRY_RE = re.compile(r"(?:\s|\.{2,})\d{1,3}$")
_CITATION_IN_TEXT_RE = re.compile(r"\b(?:BBl|AS|SR|FF|RO|RS|BU|AB|BO)\s+\d")
_NUMBERED_HEADING_RE = re.compile(r"^(\d{1,2}(?:\.\d{1,2}){0,5})\.?\s+([A-ZÄÖÜÀ-Ý][^\W\d_]{2,}.*)$")
SHORT_BLOCK = 40


def _is_heading(text: str) -> bool:
    """A numbered section heading ("3.1.2 Auswirkungen auf die Kantone",
    number dotted or <= 20) or an all-capitals title line; never a sentence."""
    if (len(text) > 120 or text.endswith((".", ":", ";", ","))
            or _TOC_ENTRY_RE.search(text) or _CITATION_IN_TEXT_RE.search(text)):
        return False
    m = _NUMBERED_HEADING_RE.match(text)
    if m:
        num = m.group(1)
        return "." in num or int(num) <= 20
    if any(ch.isdigit() for ch in text):
        return False
    words = [w for w in re.findall(r"[^\W\d_]+", text) if len(w) >= 3]
    return len(words) >= 2 and all(w.isupper() for w in words)


def _is_toc_page(blocks: list[str]) -> bool:
    entries = sum(1 for b in blocks if _TOC_ENTRY_RE.search(re.sub(r"\s+", " ", b).strip()))
    return entries >= 5


def _recurring(pages: list[list[str]], keyfn) -> set[str]:
    """Normalised keys present on at least max(3, 30 %) of the pages."""
    if len(pages) < 3:
        return set()
    seen: dict[str, int] = {}
    for blocks in pages:
        for key in {k for b in blocks for k in keyfn(b)}:
            seen[key] = seen.get(key, 0) + 1
    threshold = max(3, int(0.3 * len(pages)))
    # A key needs words: a bare number ("#", a section number on its own
    # line) recurs on every page and must never be stripped as a running head.
    return {k for k, n in seen.items()
            if n >= threshold and len(k) < 200 and len(re.findall(r"[^\W\d_]", k)) >= 3}


def _strip_running_lines(raw: str, recurring_lines: set[str]) -> str:
    """Drop running-head / footer lines at the start or end of a block
    (a heading block often begins with the document's running head)."""
    lines = raw.split("\n")
    while lines and (not lines[0].strip() or _norm_block(lines[0]) in recurring_lines):
        lines.pop(0)
    while lines and (not lines[-1].strip() or _norm_block(lines[-1]) in recurring_lines):
        lines.pop()
    return "\n".join(lines)


def segment_blocks(pages: list[list[str]], language: str = "de",
                   paginated: bool = True,
                   ocr_pages: set[int] | frozenset[int] = frozenset()) -> Iterator[dict]:
    """Paragraph dicts (para_order, page_number, section_path, article_anchor,
    text, text_length) from per-page text blocks.

    Running heads and footers are dropped: a whole block, or a block's
    leading/trailing line, whose digit-normalised text recurs on at least
    max(3, 30 %) of the pages; so are bare page numbers. Lines inside a block
    are joined with the historical pipeline's conservative de-hyphenation and
    runs of short blocks on one page (table cells) are joined with a space;
    nothing else in the wording changes, so every paragraph is quotable
    verbatim. A heading sets section_path (table-of-contents lines, which end
    in a page number, do not); an "Art. N" heading sets article_anchor until
    the next heading. A paragraph from a page in ``ocr_pages`` (1-based)
    carries text_source 'ocr' (a paragraph never spans two pages)."""
    rec_blocks = _recurring(pages, lambda b: [_norm_block(b)]) if paginated else set()
    rec_lines = (_recurring(pages, lambda b: [_norm_block(ln) for ln in b.split("\n")])
                 if paginated else set())

    section: str | None = None
    anchor: str | None = None
    order = 0
    pending: list[str] = []
    pending_page = None

    def emit(text: str, page_no):
        nonlocal order
        order += 1
        return {
            "para_order": order,
            "page_number": page_no if paginated else None,
            "section_path": section,
            "article_anchor": anchor,
            "text": text,
            "text_length": len(text),
            "text_source": TEXT_SOURCE_OCR if page_no in ocr_pages else None,
        }

    for page_no, blocks in enumerate(pages, start=1):
        toc = _is_toc_page(blocks)
        for raw in blocks:
            if paginated:
                if _norm_block(raw) in rec_blocks:
                    continue
                raw = _strip_running_lines(raw, rec_lines)
            text = _join_lines(raw, language)
            if len(text) < 3 or _PAGE_NUMBER_RE.match(text):
                continue
            first_line = raw.strip().split("\n", 1)[0].strip()
            art = ARTICLE_HEADER_RE.match(first_line) or (
                ARTICLE_HEADER_RE.match(text) if len(text) <= 120 else None)
            heading = not toc and _is_heading(text)
            short = len(text) < SHORT_BLOCK and not heading and not art
            if pending and (not short or page_no != pending_page):
                yield emit(" ".join(pending), pending_page)
                pending = []
            if short:
                pending.append(text)
                pending_page = page_no
                continue
            if heading:
                section, anchor = text, None
            if art:
                anchor = art.group(1)
            yield emit(text, page_no)
    if pending:
        yield emit(" ".join(pending), pending_page)


# ── HTTP ───────────────────────────────────────────────────────────────

class _Http:
    """urllib wrapper: identifying User-Agent, typed errors. Replaced in tests.

    ``cache_dir`` (optional) keeps every downloaded file under the sha-1 of
    its URL, so a re-segmentation or a re-built work DB never asks Fedlex
    for the same bytes twice. A cached GET makes no request at all."""

    def __init__(self, timeout: int = 180, cache_dir: str | Path | None = None):
        self.timeout = timeout
        self.cache_dir = Path(cache_dir) if cache_dir else None
        if self.cache_dir:
            self.cache_dir.mkdir(parents=True, exist_ok=True)
        self.last_from_cache = False

    def _cache_path(self, url: str) -> Path | None:
        if not self.cache_dir:
            return None
        return self.cache_dir / (hashlib.sha1(url.encode("utf-8")).hexdigest() + ".bin")

    def _req(self, url: str, method: str = "GET", headers: dict | None = None):
        import urllib.request
        h = {"User-Agent": USER_AGENT}
        h.update(headers or {})
        return urllib.request.Request(url, method=method, headers=h)

    def head_length(self, url: str) -> int | None:
        import urllib.request
        with urllib.request.urlopen(self._req(url, "HEAD"), timeout=60) as r:
            cl = r.headers.get("Content-Length")
            return int(cl) if cl and cl.isdigit() else None

    def get_prefix(self, url: str, n: int) -> bytes:
        import urllib.request
        with urllib.request.urlopen(
                self._req(url, headers={"Range": f"bytes=0-{n - 1}"}),
                timeout=self.timeout) as r:
            return r.read(n)

    def get(self, url: str, max_bytes: int) -> bytes:
        import urllib.request
        cp = self._cache_path(url)
        if cp is not None and cp.exists():
            self.last_from_cache = True
            return cp.read_bytes()
        self.last_from_cache = False
        with urllib.request.urlopen(self._req(url), timeout=self.timeout) as r:
            cl = r.headers.get("Content-Length")
            if cl and cl.isdigit() and int(cl) > max_bytes:
                raise _TooLarge(int(cl))
            data = r.read(max_bytes + 1)
            if len(data) > max_bytes:
                raise _TooLarge(len(data))
        if cp is not None:
            tmp = cp.with_suffix(".part")
            tmp.write_bytes(data)
            os.replace(tmp, cp)
        return data

    def get_to_file(self, url: str, path: str | Path, max_bytes: int) -> int:
        """Stream a file too large for memory to ``path`` (1 MiB chunks);
        returns its size. Never cached."""
        import urllib.request
        self.last_from_cache = False
        size = 0
        with urllib.request.urlopen(self._req(url), timeout=self.timeout) as r, \
                open(path, "wb") as out:
            cl = r.headers.get("Content-Length")
            if cl and cl.isdigit() and int(cl) > max_bytes:
                raise _TooLarge(int(cl), max_bytes)
            while True:
                chunk = r.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > max_bytes:
                    raise _TooLarge(size, max_bytes)
                out.write(chunk)
        return size


def _hash_file(path: str | Path) -> tuple[str, str]:
    """(sha256, sha256 of the first SAME_FILE_PREFIX_BYTES) of a file on disk."""
    full, pre = hashlib.sha256(), None
    with open(path, "rb") as fh:
        first = fh.read(SAME_FILE_PREFIX_BYTES)
        pre = hashlib.sha256(first).hexdigest()
        full.update(first)
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            full.update(chunk)
    return full.hexdigest(), pre


def _tmp_dir_for(conn: sqlite3.Connection) -> Path:
    """Directory for streamed downloads: next to the work DB (the data
    volume in production, not the small root disk)."""
    import tempfile
    for row in conn.execute("PRAGMA database_list"):
        if row[1] == "main" and row[2]:
            d = Path(row[2]).resolve().parent / "consultations_tmp"
            d.mkdir(exist_ok=True)
            return d
    return Path(tempfile.gettempdir())


class _TooLarge(Exception):
    def __init__(self, size: int, limit: int | None = None):
        super().__init__(f"{size} bytes > {limit or MAX_FILE_BYTES}")
        self.size = size


def _host_of(url: str | None) -> str:
    from urllib.parse import urlsplit
    return urlsplit(url or "").netloc.lower()


def _file_host_allowed(url: str | None) -> bool:
    return bool(url) and _host_of(url) in ALLOWED_FILE_HOSTS


_WAYBACK_SNAPSHOT_RE = re.compile(r"^(https://web\.archive\.org/web/\d{14})/")


def download_url(file_url: str) -> str:
    """The URL whose bytes are the file. A Wayback snapshot URL (the one
    shown to users) answers with the replay frame; the ``id_`` form of the
    same snapshot answers with the archived bytes as captured."""
    return _WAYBACK_SNAPSHOT_RE.sub(r"\1id_/", file_url, count=1)


# ── Fetch ──────────────────────────────────────────────────────────────

def _queue(conn: sqlite3.Connection, max_attempts: int = MAX_ATTEMPTS) -> list[sqlite3.Row]:
    prio = " ".join(f"WHEN {c} THEN {i}" for i, c in enumerate(TEXT_TYPE_PRIORITY))
    return conn.execute(
        f"""SELECT f.file_id, f.document_uri, f.language, f.file_url, f.format,
                   f.attempts, d.doc_type_code, p.year
            FROM consultation_files f
            JOIN consultation_documents d ON d.document_uri = f.document_uri
            JOIN consultation_procedures p ON p.procedure_uri = d.procedure_uri
            WHERE f.text_policy = 'text'
              AND (f.status = '{ST_PENDING}' OR (f.status = '{ST_ERROR}' AND f.attempts < ?))
            ORDER BY CASE d.doc_type_code {prio} ELSE 99 END,
                     p.year DESC, f.document_uri,
                     CASE f.language WHEN 'de' THEN 0 WHEN 'fr' THEN 1 WHEN 'it' THEN 2 ELSE 3 END""",
        (max_attempts,)).fetchall()


def _sibling_done(conn: sqlite3.Connection, document_uri: str, file_id: int):
    return conn.execute(
        f"""SELECT file_id, byte_size, prefix_sha256, sha256 FROM consultation_files
            WHERE document_uri=? AND file_id<>? AND status IN ('{ST_DONE}', '{ST_NO_TEXT}')
              AND byte_size IS NOT NULL AND prefix_sha256 IS NOT NULL""",
        (document_uri, file_id)).fetchall()


def store_text(conn: sqlite3.Connection, file_id: int, paragraphs: list[dict]) -> None:
    conn.execute("DELETE FROM consultation_paragraphs WHERE file_id=?", (file_id,))
    conn.executemany(
        """INSERT INTO consultation_paragraphs
           (file_id, para_order, page_number, section_path, article_anchor, text, text_length,
            text_source, source_member)
           VALUES (?,?,?,?,?,?,?,?,?)""",
        [(file_id, p["para_order"], p["page_number"], p["section_path"],
          p["article_anchor"], p["text"], p["text_length"], p.get("text_source"),
          p.get("source_member")) for p in paragraphs])


def _empty_pages(pages: list[list[str]]) -> set[int]:
    """1-based numbers of the pages with (almost) no text layer."""
    return {i for i, blocks in enumerate(pages, 1) if sum(len(b.strip()) for b in blocks) < 20}


def segment_parts(parts: list[TextPart], language: str) -> list[dict]:
    """Paragraphs of every part in order, numbered through; each carries
    the part's member name (ZIP member, Excel sheet) as source_member."""
    out: list[dict] = []
    for part in parts:
        for p in segment_blocks(part.pages, language, paginated=part.paginated,
                                ocr_pages=part.ocr_pages):
            p["para_order"] = len(out) + 1
            p["source_member"] = part.member
            out.append(p)
    return out


def _part_counts(parts: list[TextPart]) -> tuple[int | None, int | None]:
    """(page_count, empty_pages) over the paginated parts; None when no
    part has pages (Word, Excel)."""
    paged = [p for p in parts if p.paginated]
    if not paged:
        return None, None
    return (sum(len(p.pages) for p in paged),
            sum(len(_empty_pages(p.pages) - p.ocr_pages) for p in paged))


def _file_text_source(paragraphs: list[dict]) -> str | None:
    n_ocr = sum(1 for p in paragraphs if p.get("text_source") == TEXT_SOURCE_OCR)
    if not n_ocr:
        return None
    return TEXT_SOURCE_OCR if n_ocr == len(paragraphs) else TEXT_SOURCE_MIXED


def fetch_one(conn: sqlite3.Connection, row, http: _Http) -> str:
    """Fetch, extract and store one file. Returns the final status. Raises
    SourceUnavailable on a refusal (403/429) so the loop stops."""
    import urllib.error
    fid, url = row["file_id"], row["file_url"]
    if not _file_host_allowed(url):
        conn.execute("UPDATE consultation_files SET status=?, error=? WHERE file_id=?",
                     (ST_UNSUPPORTED, f"host not allowed: {url}", fid))
        conn.commit()
        return ST_UNSUPPORTED
    try:
        # Same bytes as another language's file of this document? Stellungnahmen
        # and many Ergebnisberichte are one PDF filed under all three languages.
        # (Fedlex only: the archive sources file one PDF per language, and a
        # HEAD + range request would be two more Wayback requests per file.)
        siblings = (_sibling_done(conn, row["document_uri"], fid)
                    if _host_of(url) == FEDLEX_FILE_HOST else [])
        if siblings:
            size = http.head_length(url)
            match = [s for s in siblings if s["byte_size"] == size]
            if match:
                prefix = http.get_prefix(url, SAME_FILE_PREFIX_BYTES)
                psha = hashlib.sha256(prefix).hexdigest()
                for s in match:
                    if s["prefix_sha256"] == psha:
                        conn.execute(
                            """UPDATE consultation_files SET status=?, same_as_file_id=?,
                               byte_size=?, prefix_sha256=?, sha256=?, attempts=attempts+1,
                               last_attempt_at=?, error=NULL, http_status=NULL,
                               ingested_at=? WHERE file_id=?""",
                            (ST_SAME_AS, s["file_id"], size, psha, s["sha256"],
                             _now(), _now(), fid))
                        conn.commit()
                        return ST_SAME_AS
        try:
            data = http.get(download_url(url), MAX_FILE_BYTES)
        except _TooLarge as e:
            if e.size > MAX_STREAM_BYTES or not hasattr(http, "get_to_file"):
                raise
            # Too large for memory: stream it to disk and read it from there.
            big_path = _tmp_dir_for(conn) / f"file-{fid}.part"
            try:
                size = http.get_to_file(download_url(url), big_path, MAX_STREAM_BYTES)
                sha, psha = _hash_file(big_path)
                return _ingest(conn, row, big_path, size, sha, psha)
            finally:
                big_path.unlink(missing_ok=True)
    except _TooLarge as e:
        conn.execute("UPDATE consultation_files SET status=?, byte_size=?, attempts=attempts+1, "
                     "last_attempt_at=?, error=? WHERE file_id=?",
                     (ST_TOO_LARGE, e.size, _now(), str(e), fid))
        conn.commit()
        return ST_TOO_LARGE
    except urllib.error.HTTPError as e:
        if e.code == 403 and _host_of(url) == WAYBACK_HOST:
            conn.execute("UPDATE consultation_files SET status=?, http_status=?, "
                         "attempts=attempts+1, last_attempt_at=?, error=? WHERE file_id=?",
                         (ST_NOT_ARCHIVED, 403, _now(), "capture excluded from replay (HTTP 403)",
                          fid))
            conn.commit()
            return ST_NOT_ARCHIVED
        if e.code in FATAL_HTTP_STATUSES:
            raise SourceUnavailable(f"GET {url} answered HTTP {e.code}", status=e.code) from e
        status = ST_NOT_FOUND if e.code in (404, 410) else ST_ERROR
        conn.execute("UPDATE consultation_files SET status=?, http_status=?, attempts=attempts+1, "
                     "last_attempt_at=?, error=? WHERE file_id=?",
                     (status, e.code, _now(), f"HTTP {e.code}", fid))
        conn.commit()
        if status == ST_ERROR:
            raise
        return status
    except (urllib.error.URLError, OSError) as e:
        conn.execute("UPDATE consultation_files SET status=?, attempts=attempts+1, "
                     "last_attempt_at=?, error=? WHERE file_id=?",
                     (ST_ERROR, _now(), f"{type(e).__name__}: {e}"[:500], fid))
        conn.commit()
        raise
    return _ingest(conn, row, data, len(data), hashlib.sha256(data).hexdigest(),
                   hashlib.sha256(data[:SAME_FILE_PREFIX_BYTES]).hexdigest())


def _final(conn: sqlite3.Connection, fid: int, status: str, size: int, sha: str, psha: str,
           error: str | None, extraction: dict | None = None) -> str:
    """Record a downloaded file that yields no text (final, no retry, not
    the host's fault: no failure streak)."""
    conn.execute("UPDATE consultation_files SET status=?, byte_size=?, sha256=?, "
                 "prefix_sha256=?, attempts=attempts+1, last_attempt_at=?, error=?, "
                 "extraction_json=? WHERE file_id=?",
                 (status, size, sha, psha, _now(), error[:500] if error else None,
                  json.dumps(extraction, ensure_ascii=False) if extraction else None, fid))
    conn.commit()
    return status


def _is_truncated_capture(url: str, src: bytes | str | Path, size: int) -> bool:
    """A Wayback copy cut short: exactly the Common Crawl record limit, or a
    PDF whose end-of-file marker is missing."""
    if _host_of(url) != WAYBACK_HOST:
        return False
    if size == WAYBACK_TRUNCATED_BYTES:
        return True
    if isinstance(src, (bytes, bytearray)):
        tail = bytes(src[-2048:])
    else:
        with open(src, "rb") as fh:
            fh.seek(max(0, size - 2048))
            tail = fh.read()
    return b"%%EOF" not in tail


def _ingest(conn: sqlite3.Connection, row, src: bytes | Path, size: int, sha: str,
            psha: str) -> str:
    """Extract and store the text of a downloaded file (bytes, or the path
    of a streamed one). Returns the final status."""
    fid, url, fmt, lang = row["file_id"], row["file_url"], row["format"], row["language"]
    if size == 0:
        return _final(conn, fid, ST_EMPTY, 0, sha, psha,
                      "the source serves a zero-byte file")
    if _host_of(url) == WAYBACK_HOST and size == WAYBACK_TRUNCATED_BYTES:
        # PyMuPDF may salvage the first pages of the cut file: never stored
        # as if they were the whole document.
        return _final(conn, fid, ST_TRUNCATED, size, sha, psha,
                      f"archived copy cut short at {size} bytes (Common Crawl record limit)")
    try:
        parts, extraction = extract_parts(src, fmt)
    except _Unsupported as e:
        return _final(conn, fid, ST_UNSUPPORTED, size, sha, psha, str(e))
    except Exception as e:  # noqa: BLE001 - corrupt file: record, keep going
        if _is_truncated_capture(url, src, size):
            return _final(conn, fid, ST_TRUNCATED, size, sha, psha,
                          f"archived copy cut short at {size} bytes: {e}")
        return _final(conn, fid, ST_PARSE_ERROR, size, sha, psha, f"parse: {e}")
    if extraction.get("format") == "pdf" and not sum(len(p.pages) for p in parts):
        if _is_truncated_capture(url, src, size):
            return _final(conn, fid, ST_TRUNCATED, size, sha, psha,
                          f"archived copy cut short at {size} bytes (no readable page)",
                          extraction)
        return _final(conn, fid, ST_PARSE_ERROR, size, sha, psha, "parse: PDF without pages",
                      extraction)
    if _host_of(url) != FEDLEX_FILE_HOST:
        # Archive sources: one trilingual PDF is often linked from each
        # language's page under its own path, and a BJ project page can list
        # the very file Fedlex holds; index the text once (any document).
        twin = conn.execute(
            f"SELECT file_id FROM consultation_files WHERE file_id<>? AND sha256=? "
            f"AND status='{ST_DONE}' ORDER BY (source IS NOT NULL), file_id LIMIT 1",
            (fid, sha)).fetchone()
        if twin is not None:
            conn.execute(
                """UPDATE consultation_files SET status=?, same_as_file_id=?, byte_size=?,
                   sha256=?, prefix_sha256=?, attempts=attempts+1, last_attempt_at=?,
                   error=NULL, http_status=NULL, ingested_at=? WHERE file_id=?""",
                (ST_SAME_AS, twin[0], size, sha, psha, _now(), _now(), fid))
            conn.commit()
            return ST_SAME_AS
    paragraphs = segment_parts(parts, lang)
    chars = sum(p["text_length"] for p in paragraphs)
    page_count, empty = _part_counts(parts)
    status = ST_DONE if chars >= 200 else ST_NO_TEXT
    with conn:
        if status == ST_DONE:
            store_text(conn, fid, paragraphs)
        else:
            conn.execute("DELETE FROM consultation_paragraphs WHERE file_id=?", (fid,))
        conn.execute(
            """UPDATE consultation_files SET status=?, byte_size=?, sha256=?, prefix_sha256=?,
               page_count=?, char_count=?, empty_pages=?, attempts=attempts+1,
               last_attempt_at=?, error=NULL, http_status=NULL, ingested_at=?,
               extraction_json=?, text_source=NULL, ocr_pages=NULL, ocr_engine=NULL,
               ocr_at=NULL
               WHERE file_id=?""",
            (status, size, sha, psha, page_count, chars, empty, _now(), _now(),
             _extraction_json(extraction, fmt), fid))
    return status


def _extraction_json(extraction: dict, declared: str | None) -> str | None:
    """Recorded when the file was not simply the PDF or Word file it was
    declared as: a converted, unpacked or re-typed file says how it was read."""
    declared = "pdf" if declared == "pdf-a" else declared
    if len(extraction) > 1 or extraction.get("format") != declared:
        return json.dumps(extraction, ensure_ascii=False)
    return None


def fetch(conn: sqlite3.Connection, *, rate_limit: float = 0.5,
          max_seconds: float | None = None, max_files: int | None = None,
          http: _Http | None = None,
          cache_dir: str | Path | None = None,
          sleep: Callable[[float], None] = time.sleep,
          clock: Callable[[], float] = time.monotonic,
          max_consecutive_failures: int = STOP_AFTER_CONSECUTIVE_FAILURES) -> dict:
    """Work the queue in priority order under the Fedlex fetch policy: stop
    at once on 429/403, after N consecutive failures, or at the time/file
    budget; back off exponentially after a failure. Every file commits on its
    own, so a stopped run resumes where it ended.

    The archive hosts (Wayback, bj.admin.ch) are paced by HOST_MIN_INTERVAL
    and stopped on their own: a refusal or a failure streak there ends that
    host's files for the run (``stopped_hosts``), never the Fedlex files."""
    http = http or _Http(cache_dir=cache_dir)
    queue = _queue(conn)
    stats: dict = {"queued": len(queue), "processed": 0, "stopped": None,
                   "by_status": {}, "errors": 0, "stopped_hosts": {}}
    t0 = clock()
    consecutive: dict[str, int] = {}
    for i, row in enumerate(queue, 1):
        if max_files is not None and stats["processed"] >= max_files:
            stats["stopped"] = "max_files"
            break
        if max_seconds is not None and clock() - t0 >= max_seconds:
            stats["stopped"] = "max_seconds"
            break
        host = _host_of(row["file_url"])
        if host in stats["stopped_hosts"]:
            continue
        # Fedlex keeps the run-wide semantics (one streak, one stop for all).
        key = host if host in HOST_MIN_INTERVAL else ""
        pause = max(rate_limit, HOST_MIN_INTERVAL.get(host, 0.0))
        failed = False
        if hasattr(http, "last_from_cache"):
            http.last_from_cache = False
        try:
            st = fetch_one(conn, row, http)
            stats["by_status"][st] = stats["by_status"].get(st, 0) + 1
            consecutive[key] = 0
        except Exception as e:  # noqa: BLE001
            failed = True
            stats["errors"] += 1
            consecutive[key] = n_fail = consecutive.get(key, 0) + 1
            code = http_status_of(e)
            log.warning("  %s (%s) failed: %s", row["document_uri"], row["language"], e)
            reason = None
            if code in FATAL_HTTP_STATUSES:
                reason = f"http_{code}"
            elif n_fail >= max_consecutive_failures:
                reason = f"{n_fail}_consecutive_failures"
            else:
                pause = max(pause, min(BACKOFF_MAX_SECONDS,
                                       BACKOFF_BASE_SECONDS * 2 ** (n_fail - 1)))
            if reason and key:
                stats["stopped_hosts"][host] = reason
                log.error("  %s stopped (%s); its files wait for the next run", host, reason)
            elif reason:
                stats["stopped"] = reason
        stats["processed"] += 1
        if stats["stopped"]:
            log.error("  stopping (%s) at %d/%d; the next run resumes", stats["stopped"],
                      i, len(queue))
            break
        if i % 50 == 0:
            log.info("  progress %d/%d %s", i, len(queue), stats["by_status"])
        if failed or not getattr(http, "last_from_cache", False):
            sleep(pause)
    stats["elapsed_seconds"] = round(clock() - t0, 1)
    set_meta(conn, "last_fetch", {"at": _now(), **stats})
    conn.commit()
    return stats


# ── OCR ────────────────────────────────────────────────────────────────
#
# Scans without a text layer, and the scanned pages inside otherwise
# born-digital files (a Stellungnahmen collection mixes typed submissions
# with signed letters), are read with Tesseract (deu/fra/ita on the
# production host) from a 300 dpi greyscale rendering by PyMuPDF. Only pages
# whose text layer is (almost) empty are read; the text layer is never
# replaced. Every paragraph from such a page carries text_source='ocr', the
# file says how many pages are OCR, with which engine and when.

OCR_DPI = 300
OCR_LANGS = {"de": "deu", "fr": "fra", "it": "ita"}
# One collection holds submissions in every language.
OCR_ALL_LANGUAGE_TYPES = frozenset({21, 15})   # Stellungnahmen, Diverses
OCR_MIN_PAGE_CHARS = 20
OCR_MAX_PAGES_PER_FILE = 3000
OCR_SCOPES = ("no_text", "partial")


def ocr_langs(language: str | None, doc_type_code: int | None) -> str:
    first = OCR_LANGS.get(language or "", "deu")
    if doc_type_code in OCR_ALL_LANGUAGE_TYPES or language not in OCR_LANGS:
        return "+".join([first] + [v for v in ("deu", "fra", "ita") if v != first])
    return first


def tesseract_image_to_text(png: bytes, langs: str) -> str:
    """Tesseract on one page image (pytesseract where installed, else the
    tesseract CLI). One thread per process: the pass runs pages in parallel."""
    os.environ.setdefault("OMP_THREAD_LIMIT", "1")
    try:
        import pytesseract
        from PIL import Image
        return pytesseract.image_to_string(Image.open(io.BytesIO(png)), lang=langs)
    except ImportError:
        import subprocess
        res = subprocess.run(["tesseract", "stdin", "stdout", "-l", langs], input=png,
                             capture_output=True, timeout=600, check=True)
        return res.stdout.decode("utf-8", errors="replace")


def tesseract_version() -> str:
    try:
        import pytesseract
        return str(pytesseract.get_tesseract_version()).split()[0]
    except Exception:  # noqa: BLE001
        import subprocess
        try:
            out = subprocess.run(["tesseract", "--version"], capture_output=True, timeout=30)
            first = (out.stdout or out.stderr).decode("utf-8", "replace").split("\n", 1)[0]
            return first.replace("tesseract", "").strip() or "unknown"
        except (OSError, subprocess.SubprocessError):
            return "unknown"


def ocr_text_blocks(text: str) -> list[str]:
    """Tesseract prints a blank line between paragraphs and one line per
    printed line: one block per paragraph (lines joined by segment_blocks)."""
    return [b.strip() for b in re.split(r"\n[ \t]*\n", text or "") if b.strip()]


class _OcrDeadline(Exception):
    """The run's time budget ended inside a file: nothing of it is stored."""


def ocr_pdf_pages(src: bytes | str | Path, page_numbers: Iterable[int], langs: str, *,
                  image_to_text: Callable[[bytes, str], str] = tesseract_image_to_text,
                  workers: int = 2, deadline: float | None = None,
                  clock: Callable[[], float] = time.monotonic) -> dict[int, list[str]]:
    """OCR blocks per 1-based page number. Pages are rendered one at a time
    in this thread (a PyMuPDF document is not thread-safe) and recognised
    by ``workers`` Tesseract processes."""
    import fitz  # PyMuPDF
    from concurrent.futures import ThreadPoolExecutor
    wanted = sorted(set(page_numbers))
    out: dict[int, list[str]] = {}
    with _open_pdf(src) as doc, ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        pending = {}
        for n in wanted:
            if deadline is not None and clock() > deadline:
                for f in pending.values():
                    f.cancel()
                raise _OcrDeadline()
            pix = doc[n - 1].get_pixmap(dpi=OCR_DPI, colorspace=fitz.csGRAY)
            pending[n] = pool.submit(image_to_text, pix.tobytes("png"), langs)
            if len(pending) >= 2 * max(1, workers):
                done_n = min(pending)
                out[done_n] = ocr_text_blocks(pending.pop(done_n).result())
        for n in sorted(pending):
            out[n] = ocr_text_blocks(pending[n].result())
    return out


def _ocr_queue(conn: sqlite3.Connection, scopes: Iterable[str], max_pages: int) -> list[sqlite3.Row]:
    prio = " ".join(f"WHEN {c} THEN {i}" for i, c in enumerate(TEXT_TYPE_PRIORITY))
    conds = []
    if "no_text" in scopes:
        conds.append(f"(f.status = '{ST_NO_TEXT}' AND f.page_count > 0)")
    if "partial" in scopes:
        conds.append(f"(f.status = '{ST_DONE}' AND f.empty_pages > 0)")
    if not conds:
        return []
    return conn.execute(
        f"""SELECT f.file_id, f.document_uri, f.language, f.file_url, f.format, f.sha256,
                   f.byte_size, f.page_count, f.empty_pages, f.status, d.doc_type_code
            FROM consultation_files f
            JOIN consultation_documents d ON d.document_uri = f.document_uri
            WHERE f.text_policy = 'text' AND f.ocr_at IS NULL
              AND ({' OR '.join(conds)})
              AND COALESCE(f.empty_pages, f.page_count) <= ?
            ORDER BY (f.status = '{ST_DONE}'), CASE d.doc_type_code {prio} ELSE 99 END,
                     COALESCE(f.empty_pages, f.page_count), f.file_id""",
        (max_pages,)).fetchall()


def ocr_one(conn: sqlite3.Connection, row, http, *, workers: int = 2,
            image_to_text: Callable[[bytes, str], str] = tesseract_image_to_text,
            engine: str | None = None, deadline: float | None = None,
            clock: Callable[[], float] = time.monotonic) -> dict:
    """OCR the text-less pages of one file and store the file's text again
    (text layer + OCR). Downloads the file (a streamed temp file above
    MAX_FILE_BYTES) and checks it is the one fetched before (sha-256): a
    file changed upstream goes back to the fetch queue instead."""
    fid, url = row["file_id"], row["file_url"]
    result = {"file_id": fid, "ocr_pages": 0}
    if not _file_host_allowed(url):
        result["skipped"] = "host not allowed"
        return result
    big_path = None
    try:
        if row["byte_size"] and row["byte_size"] > MAX_FILE_BYTES and hasattr(http, "get_to_file"):
            big_path = _tmp_dir_for(conn) / f"ocr-{fid}.part"
            http.get_to_file(download_url(url), big_path, MAX_STREAM_BYTES)
            src: bytes | Path = big_path
            sha = _hash_file(big_path)[0]
        else:
            src = http.get(download_url(url), MAX_FILE_BYTES)
            sha = hashlib.sha256(src).hexdigest()
        if row["sha256"] and sha != row["sha256"]:
            conn.execute(f"UPDATE consultation_files SET status='{ST_PENDING}', attempts=0, "
                         "error='changed upstream since the fetch; re-queued by the OCR pass' "
                         "WHERE file_id=?", (fid,))
            conn.commit()
            result["requeued"] = True
            return result
        parts, extraction = extract_parts(src, row["format"])
        langs = ocr_langs(row["language"], row["doc_type_code"])
        for part in parts:
            # the file itself, a ZIP member or a PDF attachment: any PDF part
            if not part.paginated or part.src is None:
                continue
            empty = _empty_pages(part.pages)
            if not empty:
                continue
            got = ocr_pdf_pages(part.src, empty, langs, image_to_text=image_to_text,
                                workers=workers, deadline=deadline, clock=clock)
            for n, blocks in got.items():
                if sum(len(b.strip()) for b in blocks) >= OCR_MIN_PAGE_CHARS:
                    part.pages[n - 1] = blocks
                    part.ocr_pages.add(n)
        paragraphs = segment_parts(parts, row["language"])
    finally:
        if big_path is not None:
            big_path.unlink(missing_ok=True)
    n_ocr = sum(len(p.ocr_pages) for p in parts)
    engine = engine or f"tesseract {tesseract_version()} {langs} {OCR_DPI}dpi"
    chars = sum(p["text_length"] for p in paragraphs)
    with conn:
        if n_ocr and chars >= 200:
            store_text(conn, fid, paragraphs)
            _, still_empty = _part_counts(parts)
            conn.execute(
                f"""UPDATE consultation_files SET status='{ST_DONE}', char_count=?,
                    empty_pages=?, text_source=?, ocr_pages=?, ocr_engine=?, ocr_at=?,
                    error=NULL WHERE file_id=?""",
                (chars, still_empty, _file_text_source(paragraphs), n_ocr, engine, _now(),
                 fid))
            result["status"] = ST_DONE
        else:
            # Looked at, nothing usable: recorded so the pass does not repeat it.
            note = (f"ocr: {chars} characters on {n_ocr} pages, below the 200 of a text"
                    if n_ocr else "ocr: no text recognised on the empty pages")
            conn.execute(
                "UPDATE consultation_files SET ocr_pages=0, ocr_engine=?, ocr_at=?, "
                "error=CASE WHEN status='done' THEN error ELSE ? END WHERE file_id=?",
                (engine, _now(), note, fid))
            result["status"] = row["status"]
            n_ocr = 0
    result.update(ocr_pages=n_ocr, chars=chars, extraction=extraction.get("format"))
    return result


def ocr(conn: sqlite3.Connection, *, scopes: Iterable[str] = ("no_text",),
        max_seconds: float | None = None, max_files: int | None = None,
        max_pages_per_file: int = OCR_MAX_PAGES_PER_FILE, workers: int = 2,
        rate_limit: float = 0.5, http: _Http | None = None,
        cache_dir: str | Path | None = None,
        image_to_text: Callable[[bytes, str], str] = tesseract_image_to_text,
        sleep: Callable[[float], None] = time.sleep,
        clock: Callable[[], float] = time.monotonic) -> dict:
    """The OCR pass: files without text first, then files with empty pages
    (``scopes``), the smallest first; one file at a time, each committed on
    its own. A file the budget would cut is not started again in this run
    and keeps its state. Downloads follow the fetch policy (host pacing,
    stop on 429/403)."""
    scopes = tuple(scopes)
    http = http or _Http(cache_dir=cache_dir)
    queue = _ocr_queue(conn, scopes, max_pages_per_file)
    stats: dict = {"queued": len(queue), "processed": 0, "ocr_files": 0, "ocr_pages": 0,
                   "now_done": 0, "requeued": 0, "nothing_recognised": 0, "errors": 0,
                   "stopped": None, "scopes": list(scopes)}
    engine_v = tesseract_version()
    t0 = clock()
    deadline = t0 + max_seconds if max_seconds is not None else None
    for row in queue:
        if max_files is not None and stats["processed"] >= max_files:
            stats["stopped"] = "max_files"
            break
        if deadline is not None and clock() >= deadline:
            stats["stopped"] = "max_seconds"
            break
        langs = ocr_langs(row["language"], row["doc_type_code"])
        try:
            res = ocr_one(conn, row, http, workers=workers, image_to_text=image_to_text,
                          engine=f"tesseract {engine_v} {langs} {OCR_DPI}dpi",
                          deadline=deadline, clock=clock)
        except _OcrDeadline:
            stats["stopped"] = "max_seconds"
            log.info("  OCR budget ended inside file %s; it starts again next run",
                     row["file_id"])
            break
        except SourceUnavailable as e:
            stats["stopped"] = f"http_{e.status}"
            break
        except Exception as e:  # noqa: BLE001 - one file, not the pass
            stats["errors"] += 1
            log.warning("  OCR of file %s failed: %s", row["file_id"], e)
            code = http_status_of(e)
            if code in FATAL_HTTP_STATUSES:
                stats["stopped"] = f"http_{code}"
                break
            continue
        finally:
            stats["processed"] += 1
        if res.get("requeued"):
            stats["requeued"] += 1
        elif res.get("ocr_pages"):
            stats["ocr_files"] += 1
            stats["ocr_pages"] += res["ocr_pages"]
            if row["status"] != ST_DONE and res.get("status") == ST_DONE:
                stats["now_done"] += 1
        elif not res.get("skipped"):
            stats["nothing_recognised"] += 1
        if not getattr(http, "last_from_cache", False):
            sleep(max(rate_limit, HOST_MIN_INTERVAL.get(_host_of(row["file_url"]), 0.0)))
    stats["elapsed_seconds"] = round(clock() - t0, 1)
    set_meta(conn, "last_ocr", {"at": _now(), **stats})
    conn.commit()
    return stats


# ── Re-queue (one-off, guarded) ────────────────────────────────────────

# Files recorded as final under an older extractor that the current code
# can read. Each selector is a WHERE clause on consultation_files.
REQUEUE_SELECTORS: dict[str, str] = {
    # .doc / .xlsx / .zip manifestations (never downloaded); not the
    # off-host rows, which carry an error.
    "unsupported": (f"status = '{ST_UNSUPPORTED}' AND COALESCE(error, '') NOT LIKE "
                    f"'host not allowed%' AND format IN "
                    f"({', '.join(repr(f) for f in sorted(TEXT_FORMATS))})"),
    # above the in-memory limit, within the streamed one
    "too_large": f"status = '{ST_TOO_LARGE}' AND COALESCE(byte_size, 0) <= {MAX_STREAM_BYTES}",
    # corrupt, empty, a .doc behind a PDF link, a host without python-docx
    "parse_error": f"status = '{ST_PARSE_ERROR}'",
    # a 'PDF' without a single page: an archived copy cut short
    "no_pages": f"status = '{ST_NO_TEXT}' AND COALESCE(page_count, 0) = 0",
    # still without text after OCR: PDF Portfolios (the page is a viewer
    # placeholder, the documents are attachments) read since 2026-10-21
    "no_text_after_ocr": f"status = '{ST_NO_TEXT}' AND ocr_at IS NOT NULL",
}
GUARD_UNITS = ("opencaselaw-vernehmlassungen.service",)


def _unit_busy(units: Iterable[str] = GUARD_UNITS) -> str | None:
    """The first unit that is running (any non-idle state), or None. No
    systemctl (a dev machine): None."""
    import shutil
    import subprocess
    if not shutil.which("systemctl"):
        return None
    for u in units:
        st = subprocess.run(["systemctl", "is-active", u], capture_output=True,
                            text=True).stdout.strip()
        if st in ("active", "activating", "deactivating", "reloading", "refreshing"):
            return f"{u} is {st}"
    return None


def requeue(conn: sqlite3.Connection, what: Iterable[str], *, apply: bool = False) -> dict:
    """Count (and with ``apply`` reset to pending) the files the named
    selectors match. The reset keeps the file URL and provenance, clears
    the fetch result, and is logged in meta['requeue_log']."""
    out: dict = {"apply": apply, "selected": {}}
    ids: list[int] = []
    for name in what:
        if name not in REQUEUE_SELECTORS:
            raise ValueError(f"unknown selector {name!r}; one of {sorted(REQUEUE_SELECTORS)}")
        rows = conn.execute(
            f"SELECT file_id, COALESCE(source, 'fedlex') src, format FROM consultation_files "
            f"WHERE text_policy = 'text' AND {REQUEUE_SELECTORS[name]}").fetchall()
        by: dict[str, int] = {}
        for r in rows:
            k = f"{r['src']}/{r['format']}"
            by[k] = by.get(k, 0) + 1
        out["selected"][name] = {"files": len(rows), "by_source_format": by}
        ids.extend(r["file_id"] for r in rows)
    ids = sorted(set(ids))
    out["files"] = len(ids)
    if apply and ids:
        with conn:
            for i in range(0, len(ids), 500):
                chunk = ids[i:i + 500]
                marks = ",".join("?" * len(chunk))
                conn.execute(f"DELETE FROM consultation_paragraphs WHERE file_id IN ({marks})", chunk)
                conn.execute(
                    f"""UPDATE consultation_files SET status='{ST_PENDING}', attempts=0,
                        error=NULL, http_status=NULL, byte_size=NULL, sha256=NULL, prefix_sha256=NULL,
                        same_as_file_id=NULL, page_count=NULL, char_count=NULL,
                        empty_pages=NULL, ingested_at=NULL, extraction_json=NULL,
                        text_source=NULL, ocr_pages=NULL, ocr_engine=NULL, ocr_at=NULL
                        WHERE file_id IN ({marks})""", chunk)
            prev = conn.execute("SELECT value FROM meta WHERE key='requeue_log'").fetchone()
            log_ = json.loads(prev[0]) if prev else []
            log_.append({"at": _now(), "selectors": list(what), "files": len(ids)})
            set_meta(conn, "requeue_log", log_)
    return out


# ── Publish ────────────────────────────────────────────────────────────

def coverage_counts(conn: sqlite3.Connection) -> dict:
    """The served corpus in numbers: procedures (all / from Fedlex), the
    documents with searchable text (by type), paragraphs, year range."""
    readable = f"f.status IN ('{ST_DONE}', '{ST_SAME_AS}')"
    procs, fedlex, lo, hi = conn.execute(
        "SELECT COUNT(*), SUM(source IS NULL), MIN(year), MAX(year) FROM consultation_procedures"
    ).fetchone()
    by_type = {r[0]: int(r[1]) for r in conn.execute(
        f"""SELECT d.doc_type, COUNT(DISTINCT d.document_uri) FROM consultation_documents d
            JOIN consultation_files f ON f.document_uri = d.document_uri
            WHERE {readable} GROUP BY 1""")}
    docs = conn.execute(f"SELECT COUNT(DISTINCT document_uri) FROM consultation_files f "
                        f"WHERE {readable}").fetchone()[0]
    return {
        "procedures": int(procs or 0), "procedures_fedlex": int(fedlex or 0),
        "documents": int(docs or 0), "by_doc_type": by_type,
        "paragraphs": int(conn.execute("SELECT COUNT(*) FROM consultation_paragraphs").fetchone()[0]),
        "year_range": {"from": lo, "to": hi} if lo else None,
        "counted_at": _now(),
    }


def publish(conn: sqlite3.Connection, out_path: str | Path) -> Path:
    """Atomic snapshot of the work DB onto the served path. The .tmp sits
    next to the REAL file (a symlinked output/ entry pointing at the data
    volume must be replaced through the link, on the same filesystem)."""
    out = Path(out_path)
    real = Path(os.path.realpath(str(out)))
    tmp = real.with_name(real.name + ".tmp")
    if tmp.exists():
        tmp.unlink()
    set_meta(conn, "published_at", _now())
    # Counts for stats.json (generate_stats reads them from meta instead of
    # counting ~10M paragraphs during the nightly build).
    set_meta(conn, "coverage", coverage_counts(conn))
    conn.commit()
    conn.execute("VACUUM INTO ?", (str(tmp),))
    chk = sqlite3.connect(f"file:{tmp}?mode=ro", uri=True)
    try:
        ok = chk.execute("PRAGMA quick_check").fetchone()[0]
        n = chk.execute("SELECT COUNT(*) FROM consultation_procedures").fetchone()[0]
    finally:
        chk.close()
    if ok != "ok" or n == 0:
        tmp.unlink(missing_ok=True)
        raise RuntimeError(f"snapshot failed its check (quick_check={ok!r}, procedures={n})")
    os.replace(str(tmp), str(real))
    log.info("published %s (%d procedures)", real, n)
    return real


def publish_due(conn: sqlite3.Connection, out_path: str | Path, max_age_days: int = 7) -> bool:
    """A publish without new files is still due when the served file is
    missing or older than ``max_age_days`` (procedure metadata — status,
    deadlines — changes without new files)."""
    if not Path(os.path.realpath(str(out_path))).exists():
        return True
    row = conn.execute("SELECT value FROM meta WHERE key='published_at'").fetchone()
    if not row:
        return True
    try:
        last = datetime.fromisoformat(row[0])
    except ValueError:
        return True
    return (datetime.now(timezone.utc) - last).days >= max_age_days


# ── Status ─────────────────────────────────────────────────────────────

def status_summary(conn: sqlite3.Connection) -> dict:
    out: dict = {
        "procedures": conn.execute("SELECT COUNT(*) FROM consultation_procedures").fetchone()[0],
        "documents": conn.execute("SELECT COUNT(*) FROM consultation_documents").fetchone()[0],
        "paragraphs": conn.execute("SELECT COUNT(*) FROM consultation_paragraphs").fetchone()[0],
        "files": {},
        "by_source": {},
    }
    for r in conn.execute(
            """SELECT d.doc_type, f.status, COUNT(*) n FROM consultation_files f
               JOIN consultation_documents d USING(document_uri)
               GROUP BY 1, 2 ORDER BY 1, 2"""):
        out["files"].setdefault(r["doc_type"], {})[r["status"]] = r["n"]
    for r in conn.execute(
            """SELECT COALESCE(source, 'fedlex') src, status, COUNT(*) n
               FROM consultation_files GROUP BY 1, 2 ORDER BY 1, 2"""):
        out["by_source"].setdefault(r["src"], {})[r["status"]] = r["n"]
    out["text_source"] = text_source_summary(conn)
    return out


def text_source_summary(conn: sqlite3.Connection) -> dict:
    """Files with text by where the text comes from (OCR counted apart from
    the files' own text), and what the OCR pass has left. From the file
    rows only: no scan of the paragraph table."""
    out: dict = {"files": {}, "ocr_pages": 0, "ocr_backlog": {}}
    for r in conn.execute(
            f"""SELECT COALESCE(text_source, 'document') ts, COUNT(*) n,
                       COALESCE(SUM(ocr_pages), 0) p
                FROM consultation_files WHERE status = '{ST_DONE}' GROUP BY 1"""):
        out["files"][r["ts"]] = r["n"]
        out["ocr_pages"] += r["p"]
    nt = conn.execute(
        f"""SELECT COUNT(*), COALESCE(SUM(page_count), 0) FROM consultation_files
            WHERE status = '{ST_NO_TEXT}' AND page_count > 0 AND ocr_at IS NULL""").fetchone()
    pa = conn.execute(
        f"""SELECT COUNT(*), COALESCE(SUM(empty_pages), 0) FROM consultation_files
            WHERE status = '{ST_DONE}' AND empty_pages > 0 AND ocr_at IS NULL""").fetchone()
    seen = conn.execute(
        f"""SELECT COUNT(*) FROM consultation_files
            WHERE status = '{ST_NO_TEXT}' AND ocr_at IS NOT NULL""").fetchone()
    out["ocr_backlog"] = {"no_text_files": nt[0], "no_text_pages": nt[1],
                          "partial_files": pa[0], "partial_empty_pages": pa[1]}
    out["no_text_after_ocr"] = seen[0]
    return out


def main(argv: list[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)-7s %(message)s",
                        datefmt="%H:%M:%S")
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("command", choices=["discover", "fetch", "publish", "status", "run",
                                        "ocr", "requeue"])
    ap.add_argument("--work-db", required=True)
    ap.add_argument("--out", help="served database path (publish / run)")
    ap.add_argument("--years", help="discover only these years, e.g. 2023,2024")
    ap.add_argument("--rate-limit", type=float, default=0.5)
    ap.add_argument("--max-seconds", type=float)
    ap.add_argument("--max-files", type=int)
    ap.add_argument("--cache-dir", help="keep downloaded files here (optional)")
    ap.add_argument("--scope", default="no_text",
                    help=f"ocr: comma list of {', '.join(OCR_SCOPES)} (default no_text)")
    ap.add_argument("--workers", type=int, default=2, help="ocr: Tesseract processes")
    ap.add_argument("--max-pages-per-file", type=int, default=OCR_MAX_PAGES_PER_FILE)
    ap.add_argument("--what", help=f"requeue: comma list of {', '.join(REQUEUE_SELECTORS)}")
    ap.add_argument("--apply", action="store_true",
                    help="requeue: reset the rows (default: count only)")
    ap.add_argument("--no-unit-check", action="store_true",
                    help="requeue / ocr on a COPY of the work DB: skip the check that "
                         "opencaselaw-vernehmlassungen.service is idle")
    args = ap.parse_args(argv)

    if args.command in ("requeue", "ocr") and not args.no_unit_check:
        busy = _unit_busy()
        if busy:
            log.error("refusing: %s (it writes the work DB); try again when it is idle", busy)
            return 3
    conn = connect_work_db(args.work_db)
    rc = 0
    changed = False
    if args.command in ("discover", "run"):
        years = [int(y) for y in args.years.split(",")] if args.years else None
        ds = discover(conn, years=years)
        print(json.dumps(ds, ensure_ascii=False))
        changed |= bool(ds["files_new"] or ds["files_requeued"])
    if args.command in ("fetch", "run"):
        st = fetch(conn, rate_limit=args.rate_limit, max_seconds=args.max_seconds,
                   max_files=args.max_files, cache_dir=args.cache_dir)
        print(json.dumps(st, ensure_ascii=False))
        changed |= bool(st["processed"])
        if st["stopped"] and st["stopped"] not in ("max_seconds", "max_files"):
            rc = INGEST_STOPPED_EXIT_CODE
    if args.command in ("publish", "run"):
        if not args.out:
            ap.error("--out is required for publish")
        if args.command == "publish" or changed or publish_due(conn, args.out):
            publish(conn, args.out)
        else:
            log.info("nothing new since the last publish; %s left as it is", args.out)
    if args.command == "status":
        print(json.dumps(status_summary(conn), ensure_ascii=False, indent=2))
    if args.command == "requeue":
        if not args.what:
            ap.error("--what is required for requeue")
        print(json.dumps(requeue(conn, [w.strip() for w in args.what.split(",") if w.strip()],
                                 apply=args.apply), ensure_ascii=False, indent=2))
    if args.command == "ocr":
        scopes = [x.strip() for x in args.scope.split(",") if x.strip()]
        bad = [x for x in scopes if x not in OCR_SCOPES]
        if bad:
            ap.error(f"unknown --scope {bad}; one of {OCR_SCOPES}")
        st = ocr(conn, scopes=scopes, max_seconds=args.max_seconds, max_files=args.max_files,
                 max_pages_per_file=args.max_pages_per_file, workers=args.workers,
                 rate_limit=args.rate_limit, cache_dir=args.cache_dir)
        print(json.dumps(st, ensure_ascii=False))
        if st["stopped"] and st["stopped"].startswith("http_"):
            rc = INGEST_STOPPED_EXIT_CODE
    conn.close()
    return rc


if __name__ == "__main__":
    sys.exit(main())
